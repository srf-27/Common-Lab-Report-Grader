import os
import re
from zipfile import ZipFile
import pymupdf

from .. import storage
from .doc_reader import read_doc


ARCHIVE_EXTENSIONS = {
    ".zip": "zip",
    ".rar": "rar",
    ".7z": "7z",
}

ARCHIVE_SIGNATURES = {
    b"PK\x03\x04": "zip",
    b"PK\x05\x06": "zip",
    b"Rar!\x1a\x07": "rar",
    b"7z\xbc\xaf\x27\x1c": "7z",
}

REPORT_EXTENSIONS = {".pdf", ".docx", ".doc", ".md"}

# 本地解析器不处理、可由 markitdown MCP 兜底转换的格式
MARKITDOWN_EXTENSIONS = {".html", ".htm", ".pptx", ".xlsx", ".epub"}

COLLECT_EXTENSIONS = REPORT_EXTENSIONS | MARKITDOWN_EXTENSIONS


def detect_archive_type(filepath):
    """按扩展名判断压缩包类型，扩展名不可信时按文件头兜底，其他返回 None。"""
    fmt = ARCHIVE_EXTENSIONS.get(os.path.splitext(filepath)[1].lower())
    if fmt:
        return fmt
    with open(filepath, "rb") as f:
        header = f.read(8)
    for signature, fmt in ARCHIVE_SIGNATURES.items():
        if header.startswith(signature):
            return fmt
    return None


def extract_archive(filepath, dest_dir):
    fmt = detect_archive_type(filepath)
    if fmt is None:
        raise ValueError(f"Not a recognized archive: {filepath}")
    os.makedirs(dest_dir, exist_ok=True)
    if fmt == "zip":
        with ZipFile(filepath, "r") as zf:
            zf.extractall(dest_dir)
    elif fmt == "rar":
        import rarfile
        with rarfile.RarFile(filepath, "r") as rf:
            rf.extractall(dest_dir)
    elif fmt == "7z":
        import py7zr
        with py7zr.SevenZipFile(filepath, "r") as sz:
            sz.extractall(dest_dir)
    return dest_dir


def extract_pdf(filepath, output_dir, ocr_engine=None):
    os.makedirs(output_dir, exist_ok=True)
    os.makedirs(os.path.join(output_dir, "images"), exist_ok=True)

    doc = pymupdf.open(filepath)
    pages = _pdf_pages(doc)
    full_text = []
    image_records = []
    seen_xrefs = set()
    pages_without_text = 0
    page_count = len(doc)

    for page_num in range(page_count):
        page = doc[page_num]
        text = pages[page_num]
        if page.get_text("text").strip():
            full_text.append(f"<!-- Page {page_num + 1} -->\n{text.strip()}")
        else:
            pages_without_text += 1
            if ocr_engine:
                pix = page.get_pixmap(dpi=300)
                ocr_text = ocr_engine.recognize_bytes(pix.tobytes("png"))
                full_text.append(f"<!-- Page {page_num + 1} (OCR) -->\n{ocr_text}")

        img_index = 0
        for img_info in page.get_images(full=True):
            xref = img_info[0]
            if xref in seen_xrefs:
                continue
            seen_xrefs.add(xref)
            base_image = doc.extract_image(xref)
            if base_image:
                ext = base_image["ext"]
                context_label = _get_image_context(text, img_index)
                img_name = f"fig_{page_num+1:03d}_{img_index+1:03d}_{context_label}.{ext}"
                img_name = _sanitize_filename(img_name)
                img_path = os.path.join(output_dir, "images", img_name)
                with open(img_path, "wb") as f:
                    f.write(base_image["image"])
                image_records.append({
                    "filename": os.path.basename(img_path),
                    "page": page_num + 1,
                    "context": context_label,
                })
                img_index += 1

    text_path = os.path.join(output_dir, "text.md")
    storage.write_text(text_path, "\n\n".join(full_text))

    doc.close()

    meta = {
        "filename": os.path.basename(filepath),
        "pages": page_count,
        "image_count": len(image_records),
        "is_scanned": pages_without_text > 0,
    }
    return meta, image_records


def _pdf_pages(doc):
    """逐页正文：装了 pymupdf4llm 就用它的版面分析（标题层级与阅读顺序更准）。"""
    try:
        import pymupdf4llm

        chunks = pymupdf4llm.to_markdown(doc, page_chunks=True)
    except Exception:
        return [_page_body(page) for page in doc]
    markdown = {}
    for chunk in chunks:
        metadata = chunk.get("metadata") or {}
        number = metadata.get("page") or metadata.get("page_number")
        if number:
            markdown[number - 1] = _strip_image_refs(chunk.get("text") or "")
    return [markdown.get(index) or _page_body(page) for index, page in enumerate(doc)]


_PDF_IMAGE_REF = re.compile(r"!\[\]\([^)]*\)\n?")


def _strip_image_refs(text):
    return _PDF_IMAGE_REF.sub("", text)


def _page_body(page):
    """无 pymupdf4llm 时：文本块与识别到的表格按纵向位置合并，表格区域不重复出文本。"""
    tables = page.find_tables().tables
    blocks = page.get_text("blocks")
    if tables:
        blocks = [b for b in blocks
                  if not any(_center_inside(b[:4], table.bbox) for table in tables)]
    parts = [(b[1], b[4].strip()) for b in blocks if b[4].strip()]
    parts += [(table.bbox[1], table.to_markdown()) for table in tables]
    parts.sort(key=lambda part: part[0])
    return "\n\n".join(part[1] for part in parts)


def _center_inside(block, box):
    x = (block[0] + block[2]) / 2
    y = (block[1] + block[3]) / 2
    return box[0] <= x <= box[2] and box[1] <= y <= box[3]


def extract_docx(filepath, output_dir):
    os.makedirs(output_dir, exist_ok=True)
    os.makedirs(os.path.join(output_dir, "images"), exist_ok=True)

    from docx import Document
    from docx.table import Table

    doc = Document(filepath)
    full_text = [_table_to_markdown(block) if isinstance(block, Table) else block.text
                 for block in doc.iter_inner_content()]
    image_records = []
    with ZipFile(filepath, "r") as docx_zip:
        for index, (media, context) in enumerate(_docx_media_in_order(doc), 1):
            ext = os.path.splitext(media)[1].lstrip(".")
            label = _sanitize_filename(context)[:30] or "no_context"
            img_name = _sanitize_filename(f"fig_{index:03d}_{label}.{ext}")
            img_path = os.path.join(output_dir, "images", img_name)
            with docx_zip.open(f"word/media/{media}") as src, open(img_path, "wb") as dst:
                dst.write(src.read())
            image_records.append({
                "filename": os.path.basename(img_path),
                "context": context[:40] or "no_context",
            })

    text_path = os.path.join(output_dir, "text.md")
    storage.write_text(text_path, "\n\n".join(full_text))

    meta = {
        "filename": os.path.basename(filepath),
        "paragraphs": len(doc.paragraphs),
        "image_count": len(image_records),
        "is_scanned": False,
    }
    return meta, image_records


def _docx_media_in_order(doc):
    """按正文出现顺序返回 [(媒体文件名, 所在段落文本)]，同一张图只记一次。"""
    from docx.oxml.ns import qn

    placed = []
    last_text = ""
    for para in doc.element.body.iter(qn("w:p")):
        text = "".join(node.text or "" for node in para.iter(qn("w:t"))).strip()
        if text:
            last_text = text
        for blip in para.iter(qn("a:blip")):
            rel = doc.part.rels.get(blip.get(qn("r:embed")) or "")
            if rel is None or rel.is_external:
                continue
            media = os.path.basename(str(rel.target_part.partname))
            if any(media == name for name, _ in placed):
                continue
            placed.append((media, text or last_text))
    return placed


def _table_to_markdown(table):
    rows = [[cell.text.strip().replace("|", "\\|").replace("\n", " ") for cell in row.cells]
            for row in table.rows]
    if not rows:
        return ""
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |",
             "| " + " | ".join(["---"] * len(header)) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def extract_markdown(filepath, output_dir):
    """Markdown 报告：正文原样保留，![]() 引用的本地图片一并收进来。"""
    image_records, text = _write_linked_images(filepath, output_dir, _read_text(filepath))
    meta, _ = _write_text(filepath, output_dir, text, image_count=len(image_records))
    return meta, image_records


_MARKDOWN_IMAGE = re.compile(r"!\[([^\]]*)\]\(\s*<?([^)>\s]+)>?\s*\)")


def _write_linked_images(filepath, output_dir, text):
    base_dir = os.path.dirname(os.path.abspath(filepath))
    records = []

    def replace(match):
        label, ref = match.group(1), match.group(2)
        if "://" in ref or ref.startswith("data:"):
            return match.group(0)
        source = os.path.abspath(os.path.join(base_dir, ref))
        if not source.startswith(base_dir + os.sep) or not os.path.isfile(source):
            return match.group(0)
        ext = os.path.splitext(source)[1].lower()
        stem = label or os.path.splitext(os.path.basename(source))[0]
        name = _sanitize_filename(f"fig_{len(records) + 1:03d}_{stem}") + ext
        os.makedirs(os.path.join(output_dir, "images"), exist_ok=True)
        with open(source, "rb") as src, open(os.path.join(output_dir, "images", name), "wb") as dst:
            dst.write(src.read())
        records.append({"filename": name, "context": label or os.path.basename(source)})
        return f"![{label}](images/{name})"

    return records, _MARKDOWN_IMAGE.sub(replace, text)


def extract_doc(filepath, output_dir):
    """老式 .doc：取正文与内嵌图片；若是 zip 型（docx 改名）则按 docx 处理。"""
    with open(filepath, "rb") as f:
        if f.read(2) == b"PK":
            return extract_docx(filepath, output_dir)
    text, images = read_doc(filepath)
    image_records = _write_images(output_dir, images)
    meta, _ = _write_text(filepath, output_dir, text, image_count=len(image_records))
    return meta, image_records


def _write_images(output_dir, images):
    if not images:
        return []
    os.makedirs(os.path.join(output_dir, "images"), exist_ok=True)
    records = []
    for index, (ext, blob) in enumerate(images, 1):
        name = _sanitize_filename(f"fig_{index:03d}.{ext}")
        with open(os.path.join(output_dir, "images", name), "wb") as f:
            f.write(blob)
        records.append({"filename": name, "context": "no_context"})
    return records


def _write_text(filepath, output_dir, text, image_count=0):
    os.makedirs(output_dir, exist_ok=True)
    storage.write_text(os.path.join(output_dir, "text.md"), text)
    meta = {
        "filename": os.path.basename(filepath),
        "chars": len(text),
        "image_count": image_count,
        "is_scanned": False,
    }
    return meta, []


def _read_text(filepath):
    raw = open(filepath, "rb").read()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _get_image_context(page_text, img_index):
    lines = [l.strip() for l in page_text.split("\n") if l.strip()]
    heading_keywords = ["图", "Fig", "Figure", "表", "Table", "实验", "结果", "数据", "装置"]
    for line in reversed(lines):
        if any(kw in line for kw in heading_keywords):
            return line[:40]
    if lines:
        return lines[-1][:40]
    return "no_context"


def _sanitize_filename(name):
    keepchars = (" ", ".", "_", "-")
    return "".join(c for c in name if c.isalnum() or c in keepchars).rstrip()


def collect_report_files(directory):
    supported = {*COLLECT_EXTENSIONS, *ARCHIVE_EXTENSIONS}
    files = []
    for root, _, filenames in os.walk(directory):
        for fn in filenames:
            if os.path.splitext(fn)[1].lower() in supported:
                files.append(os.path.join(root, fn))
    return files
