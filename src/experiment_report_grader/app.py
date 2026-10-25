"""Web 路由层：会话 → 细则 → 报告 → 评分 → 终审 → 导出，只做编排，逻辑落在 nodes/utils/web。"""

import csv
import hashlib
import io
import json
import os
import re
import threading
import time
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait

import yaml
from urllib.parse import quote
from flask import Flask, jsonify, redirect, request, send_file

from .embedding import Chroma, create_embedding_function
from .factory import llm
from .mcp import client as mcp_client
from .mcp import store as mcp_store
from .nodes import agreement, external_scores, rubric_io
from .nodes.grader import Grader, _coerce_rubric
from .nodes.report_extractor import ReportExtractor
from .nodes.rubric_generator import RubricGenerator
from .utils.file_parser import collect_report_files
from . import storage
from .web.state import (
    OUTPUT_FOLDER, PROJECT_ROOT, RUNNING_STATUS, SESSIONS_DIR, STUDENT_FIELDS, VALID_LEVELS,
    hydrate_session, load_mcp_mount, load_roster, load_sessions, new_session,
    remove_session, sanitize_filename, save_approved, save_mcp_mount, save_roster,
    save_sessions, session_dir, session_sub, state,
)
from .web import tasks as task_mod
from .web.tasks import task_cancelled, task_finish, task_running, task_start, task_update

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 500 * 1024 * 1024
app.config["OUTPUT_FOLDER"] = OUTPUT_FOLDER
app.config["SESSIONS_DIR"] = SESSIONS_DIR
# 绝对路径推导：__file__ 相对/绝对随启动方式变化，dirname×4 会落到项目外层目录
CONFIG_PATH = os.path.join(PROJECT_ROOT, "config", "config.yaml")

STAGES = [
    ("initial", "① 初始（未配置细则）"),
    ("rubric", "② 细则迭代"),
    ("grading", "③ 评分中"),
    ("review", "④ 待终审"),
    ("done", "⑤ 完成"),
]


def _bootstrap():
    idx = load_sessions()
    if not idx["sessions"]:
        new_session("默认课程")
    else:
        hydrate_session(idx["current"] or idx["sessions"][-1]["id"])


storage.configure(state["config"])  # 配了 PG 则全量入库；连不上静默回退文件
_bootstrap()


# ---------------------------------------------------------------- 通用小工具
def _read_text(path):
    return storage.read_text(path)


def _read_doc_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        import pymupdf
        doc = pymupdf.open(filepath)
        try:
            return "\n".join(page.get_text("text") for page in doc)
        finally:
            doc.close()
    if ext == ".docx":
        from docx import Document
        return "\n".join(p.text for p in Document(filepath).paragraphs)
    return _read_text(filepath)


def _chunks(text, limit=30):
    r"""按空行切段，但跨行的 \[ ... \] 公式块保持在同一块内。"""
    blocks = [b for b in re.split(r"\n\s*\n", text) if b.strip()]
    merged, pending = [], None
    for block in blocks:
        if pending is not None:
            pending += "\n\n" + block
            if "\\]" in block:
                merged.append(pending)
                pending = None
            continue
        if "\\[" in block and "\\]" not in block:
            pending = block
        else:
            merged.append(block)
    if pending:
        merged.append(pending)
    return merged[:limit]


def _report_text(entry):
    return _read_text(os.path.join(entry["meta"]["output_dir"], "text.md"))


def _graded_ids():
    return {r["report_id"] for r in (state["grader"].results if state["grader"] else [])}


def _approved_ids():
    return state.get("approved_ids") or set()


def _rubric_mtime():
    p = os.path.join(session_dir(), "rubrics", "rubric_current.json")
    return os.path.getmtime(p) if os.path.exists(p) else None


def _result_stale(result):
    if result.get("status") not in ("graded", "llm_failed", "flagged_injection", "blocked_safety"):
        return False
    graded_at = result.get("graded_at")
    mtime = _rubric_mtime()
    return bool(graded_at and mtime and graded_at < mtime)


def _student_fields():
    roster = load_roster()
    return roster["headers"] if roster else STUDENT_FIELDS


def _find_entry(filename):
    return next((r for r in state["extracted_reports"]
                 if r.get("meta", {}).get("filename") == filename), None)


def _origin_path(filename):
    name = os.path.basename(filename)
    for root in (session_sub("uploads"), session_sub("reports")):
        for cur, _, files in os.walk(root):
            if name in files:
                return os.path.join(cur, name)
    return None


def _download_paths():
    """评分产物相对会话目录的下载路径（文件存在才返回）。"""
    out = {}
    for key, name in (("excel", "grading_summary.xlsx"), ("json", "grading_results.json"),
                      ("flagged", "flagged_students.xlsx")):
        path = os.path.join(session_sub("grading"), name)
        if os.path.exists(path):
            out[key] = os.path.relpath(path, session_dir()).replace("\\", "/")
    return out


def _busy():
    return task_running() or (state.get("grading_status") or {}).get("phase") == "running"


# ---------------- 会话阶段
def _session_stage():
    """纯数据推导当前会话阶段，返回 (key, label)。"""
    if not state.get("rubric_json") and not state.get("rubric"):
        return STAGES[0]
    results = state["grader"].results if state["grader"] else []
    stale = any(_result_stale(r) for r in results)
    running = (state.get("grading_status") or {}).get("phase") == "running"
    if not results or stale:
        label = STAGES[1][1] + ("（评分结果已过期，请重新评分）" if stale else "")
        if not stale and state.get("rubric_json") and not _rubric_confirmed():
            label += "（待样本确认）"
        return STAGES[2] if running else ("rubric", label)
    graded = [r for r in results if r.get("status") == "graded"]
    total = len([r for r in state["extracted_reports"] if r.get("status") == "ok"])
    unapproved = [r for r in graded if r["report_id"] not in _approved_ids()]
    if running or len(graded) < total:
        return STAGES[2]
    if unapproved or len(results) > len(graded):
        return STAGES[3]
    return STAGES[4]


# ---------------------------------------------------------------- 会话
@app.route("/session/new", methods=["POST"])
def session_new():
    hydrate_session(new_session(request.form.get("name", ""))["id"])
    return redirect("/app/workbench")


@app.route("/session/switch", methods=["POST"])
def session_switch():
    sid = request.form.get("session_id", "")
    idx = load_sessions()
    if any(s["id"] == sid for s in idx["sessions"]):
        idx["current"] = sid
        save_sessions(idx)
        hydrate_session(sid)
    return redirect(request.form.get("next") or "/app/workbench")


@app.route("/session/remove", methods=["POST"])
def session_remove():
    sid = request.form.get("session_id", "")
    RUNNING_STATUS.pop(sid, None)
    task_mod.RUNNING_TASKS.pop(sid, None)
    entry = remove_session(sid)
    hydrate_session(entry["id"])
    return redirect("/app/workbench")


# ---------------------------------------------------------------- HeroUI 预览（React islands）
UI_DIST = os.path.join(os.path.dirname(__file__), "static", "ui")


def _bootstrap_payload():
    idx = load_sessions()
    current = next((s for s in idx["sessions"] if s["id"] == state.get("session")), {})
    results = state["grader"].results if state.get("grader") else []
    ok_reports = [r for r in state["extracted_reports"] if r.get("status") == "ok"]
    graded = {r["report_id"] for r in results if r.get("status") == "graded"}
    approved = _approved_ids()
    uploads = os.path.join(session_dir(), "uploads")
    names = os.listdir(uploads) if os.path.isdir(uploads) else []
    stage_key, stage_label = _session_stage()
    agreement = _agreement()
    return {
        "session": {"id": state.get("session"), "name": current.get("name", "")},
        "sessions": idx["sessions"],
        "stage": {"key": stage_key, "label": stage_label},
        "rubric": {
            "count": len((state.get("rubric_json") or {}).get("criteria") or []),
            "confirmed": _rubric_confirmed(),
            "hasUpload": any(os.path.splitext(n)[1].lower() in (".docx", ".pdf", ".txt")
                             and not n.startswith("calibration_") for n in names),
        },
        "reports": {
            "total": len(ok_reports),
            "graded": len(graded & {r["meta"]["filename"] for r in ok_reports}),
            "hasUpload": any(os.path.splitext(n)[1].lower() in
                             (".pdf", ".docx", ".doc", ".md", ".zip", ".rar", ".7z",
                              ".html", ".htm", ".pptx", ".xlsx", ".epub")
                             and not n.startswith("calibration_") for n in names),
        },
        "approved": len(approved),
        "unapproved": len([r for r in results if r["report_id"] not in approved]),
        "mcp": {"mounted": len((state.get("mcp") or {}).get("servers") or [])},
        "usage": state.get("usage") or {},
        "task": state.get("task"),
        "agreement": ({"n": agreement["overall"]["n"],
                       "cohen_kappa": agreement["overall"]["cohen_kappa"],
                       "weighted_kappa": agreement["overall"]["weighted_kappa"],
                       "krippendorff_alpha": agreement["overall"]["krippendorff_alpha"]}
                      if agreement and agreement["overall"]["n"] else None),
    }


def _agreement():
    """仅在存在教师真值（teacher_grades.json）时计算一致性，避免无谓开销。"""
    from .nodes.agreement import agreement_report, load_teacher_grades
    if not state["grader"]:
        return None
    teacher = load_teacher_grades(session_dir())
    if not teacher:
        return None
    return agreement_report(state["grader"].results, state.get("rubric_json") or {}, teacher)


# ---------------------------------------------------------------- 细则
@app.route("/upload_criteria", methods=["POST"])
def upload_criteria():
    f = request.files.get("criteria_file")
    if not f:
        return jsonify({"error": "未收到文件"}), 400
    if _busy():
        return jsonify({"error": "有任务正在进行，请稍候或先取消"}), 409
    name = sanitize_filename(f.filename)
    filepath = os.path.join(session_sub("uploads"), "criteria_" + name)
    f.save(filepath)
    state["criteria_upload"] = filepath

    experiment = request.files.get("experiment_file")
    experiment_name = ""
    if experiment and experiment.filename:
        experiment_name = sanitize_filename(experiment.filename)
        experiment_path = os.path.join(session_sub("uploads"), "experiment_" + experiment_name)
        experiment.save(experiment_path)
        state["criteria_experiment"] = experiment_path
    else:
        state["criteria_experiment"] = None
    return jsonify({"ok": True, "filename": name, "experiment": experiment_name})


@app.route("/criteria/clear", methods=["POST"])
def criteria_clear():
    """删除已上传的标准/实验流程文件（不影响已生成的细则本身）。"""
    upload_dir = session_sub("uploads")
    removed = 0
    for name in os.listdir(upload_dir):
        if name.startswith(("criteria_", "experiment_")):
            try:
                os.remove(os.path.join(upload_dir, name))
                removed += 1
            except OSError:
                pass
    state.pop("criteria_upload", None)
    state.pop("criteria_experiment", None)
    return jsonify({"ok": True, "removed": removed})


class _ThinkingTap:
    """LLM 客户端包装：流式调用的思考/生成增量实时写入任务进度，其余行为不变。"""

    def __init__(self, client, task):
        self._client = client
        self._task = task
        self._saw_thinking = False

    def chat(self, messages, temperature=0.3, retries=2):
        answer = []
        try:
            for kind, payload in self._client.chat_stream(messages, temperature=temperature):
                if kind == "thinking" and payload:
                    self._saw_thinking = True
                    task_update(self._task, thinking=payload)
                elif kind == "answer" and payload:
                    answer.append(payload)
                    if not self._saw_thinking:
                        task_update(self._task, thinking=payload)
            return "".join(answer)
        except Exception:
            return self._client.chat(messages, temperature=temperature, retries=retries)


@app.route("/process_criteria", methods=["POST"])
def process_criteria():
    filepath = state.get("criteria_upload")
    if not filepath or not os.path.exists(filepath):
        updir = session_sub("uploads")
        candidates = [os.path.join(updir, n) for n in os.listdir(updir)
                      if os.path.splitext(n)[1].lower() in (".docx", ".pdf", ".txt")]
        if not candidates:
            return redirect("/app/workbench")
        filepath = max(candidates, key=os.path.getmtime)
        state["criteria_upload"] = filepath
    if _busy():
        return redirect("/app/workbench")
    if request.form.get("confirm_reset") != "1" and _has_downstream():
        return jsonify({"needReset": True,
                        "error": "重新生成细则会清除已提取的报告、评分结果与向量库（上传原件保留），请确认后重试"}), 409
    # 请求线程内捕获会话与路径，线程绝不读 state / 不调 session_dir()
    sid = state["session"]
    session_path = session_dir()
    experiment_path = state.get("criteria_experiment")
    rubrics_dir = os.path.join(session_path, "rubrics")
    os.makedirs(rubrics_dir, exist_ok=True)
    t = task_start("criteria", "/app/rubric", mode="stage")
    task_update(t, detail=f"已上传 {os.path.basename(filepath)}")
    generator = RubricGenerator(llm_client=_ThinkingTap(llm(state["config"]), t),
                                output_dir=rubrics_dir)

    def work():
        try:
            task_update(t, stage="读取评分标准文本")
            doc_text = _read_doc_text(filepath)
            if not doc_text.strip():
                task_finish(t, "error", "评分标准文档没有可读文本（可能是纯扫描件，请先 OCR）")
                return
            experiment_text = ""
            if experiment_path and os.path.exists(experiment_path):
                task_update(t, stage="读取实验流程文本")
                experiment_text = _read_doc_text(experiment_path)
            task_update(t, stage="生成结构化评分细则",
                        detail="单次长调用，通常需要 1-3 分钟；取消将在当前步骤结束后生效")
            if experiment_text.strip():
                rubric_json = generator.generate_structured_pair(doc_text, experiment_text)
            else:
                rubric_json = generator.generate_structured(doc_text)
            if task_cancelled(t):
                task_finish(t, "cancelled")
                return
            warnings = rubric_json.pop("warnings", [])
            task_update(t, stage="保存评分细则")
            # 磁盘写入无条件进行（落捕获的会话目录）；内存只在仍是当前会话时更新
            # _save_rubric 内部会重置下游并作废校准样本的模型判级
            _save_rubric(rubric_json, doc_text, experiment_text, base=session_path)
            if state.get("session") == sid:
                state["criteria_doc"] = doc_text
                state["experiment_doc"] = experiment_text
                state["rubric_json"] = rubric_json
                state["rubric"] = rubric_io.criteria_to_markdown(rubric_json)
                state["rubric_warnings"] = warnings
            task_update(t, stage="完成",
                        detail=f"共 {len(rubric_json['criteria'])} 个维度"
                               + (f"；{warnings[0]['message']}" if warnings else ""))
            task_finish(t, "done")
        except Exception as e:
            task_finish(t, "cancelled" if task_cancelled(t) else "error", e)

    threading.Thread(target=work, daemon=True).start()
    return redirect("/app/workbench")


def _sync_dimension_collections(names):
    """只在 rubric 确认时调用：补齐当前维度的 collection，删掉已移除维度的旧 collection。"""
    keep = {str(name) for name in names if str(name).strip()}
    try:
        vector_cfg = (state.get("config") or {}).get("vector_store") or {}
        store = Chroma(config=vector_cfg, session_dir=session_dir(), create=False)
        existing = set()
        for collection in store.client.list_collections():
            if not collection.name.startswith("dim_"):
                continue
            dimension = (collection.metadata or {}).get("dimension")
            if not dimension:
                continue
            if dimension not in keep:
                store.client.delete_collection(collection.name)
            else:
                existing.add(dimension)
        embedding_function = create_embedding_function(vector_cfg.get("embedding"))
        for dimension in keep - existing:
            store.client.get_or_create_collection(
                name="dim_" + hashlib.sha1(dimension.encode("utf-8")).hexdigest()[:12],
                metadata={"hnsw:space": "cosine", "dimension": dimension},
                embedding_function=embedding_function)
    except Exception:
        pass


def _save_rubric(rubric_json, doc_text=None, experiment_text=None, base=None):
    """保存细则并重置下游：细则一变，提取/评分/向量/终审全部作废清除。

    保留上传原件（uploads/，重新点「处理」即可）与课外成绩单。
    base 显式传入时作用于该会话目录（后台线程用），否则当前会话。
    """
    base = base or session_dir()
    rubrics_dir = os.path.join(base, "rubrics")
    os.makedirs(rubrics_dir, exist_ok=True)
    rubric_io.save_rubric_json(rubric_json, os.path.join(rubrics_dir, "rubric_current.json"))
    storage.write_text(os.path.join(rubrics_dir, "rubric_latest.md"),
                       rubric_io.criteria_to_markdown(rubric_json))
    if doc_text is not None:
        storage.write_text(os.path.join(rubrics_dir, "criteria_doc.txt"), doc_text)
    if experiment_text is not None:
        experiment_path = os.path.join(rubrics_dir, "experiment_doc.txt")
        if experiment_text.strip():
            storage.write_text(experiment_path, experiment_text)
        else:
            storage.remove(experiment_path)
    # 校准样本的模型判级基于旧细则，一并作废（教师手填等级保留）
    if storage.exists(_pilot_path(base)):
        pilot = _load_pilot(base)
        for entry in pilot["reports"]:
            entry["ai"] = {}
            entry["status"] = "pending"
        _save_pilot(pilot, base)
    _reset_downstream(base)


def _has_downstream():
    """是否已有基于当前细则产生的数据——保存/重新生成细则将清除它们，需用户二次确认。"""
    if state.get("extracted_reports"):
        return True
    if state.get("grader") and state["grader"].results:
        return True
    if _approved_ids() or state.get("corrections"):
        return True
    try:
        store = Chroma(config=(state.get("config") or {}).get("vector_store") or {},
                       session_dir=session_dir(), create=False)
        return any(collection.count() for collection in store.client.list_collections())
    except Exception:
        return False


def _clear_session_vectors(base):
    """删除会话的全部 Chroma collections（与「向量库」页清空同款逻辑）。"""
    try:
        store = Chroma(config=(state.get("config") or {}).get("vector_store") or {},
                       session_dir=base, create=False)
        for collection in store.client.list_collections():
            store.client.delete_collection(collection.name)
    except Exception:
        pass


def _reset_downstream(base):
    """细则已变：清除评分结果、提取报告、向量库、终审/重评记录与校准工作区。"""
    sid = os.path.basename(base)
    _clear_session_vectors(base)
    if state.get("session") == sid:
        if state.get("grader") is not None:
            state["grader"].invalidate_vector_store()
        state["grader"] = None
        state["extracted_reports"] = []
        state["report_uploads"] = []
        state["approved_ids"] = set()
        state["corrections"] = {}
        state["calibration"] = None
    storage.write_json(os.path.join(base, "approved.json"), [])
    storage.write_json(os.path.join(base, "corrections.json"), {})
    storage.remove(os.path.join(base, "calibration_workspace.json"))
    for sub in ("grading", "reports", "kb_db", "calibration_kb"):
        storage.remove_tree(os.path.join(base, sub))


def _apply_external():
    """把成绩单合并进当前结果（幂等），返回写入条数。"""
    grader = state.get("grader")
    rubric = state.get("rubric_json")
    external = state.get("external") or {}
    if not grader or not rubric or not external:
        return 0
    students = {r["meta"]["filename"]: (r["meta"].get("student") or {})
                for r in state["extracted_reports"]}
    return external_scores.apply_to_results(grader.results, rubric, external, students.get)


def _read_sheet(storage):
    ext = os.path.splitext(storage.filename)[1].lower()
    if ext in (".xlsx", ".xls"):
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(storage.read()), read_only=True, data_only=True)
        rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
        workbook.close()
        return rows
    if ext == ".csv":
        text = storage.read().decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    raise ValueError("成绩单仅支持 .xlsx / .csv")


def _score_value(cell):
    if cell is None:
        return None
    text = str(cell).strip().upper()
    if text in rubric_io.LEVEL_POINTS:
        return rubric_io.LEVEL_POINTS[text]
    try:
        value = float(text.replace("分", "").strip())
    except ValueError:
        return None
    return value if 0 <= value <= 100 else None


@app.route("/rubric_external_upload", methods=["POST"])
def rubric_external_upload():
    upload = request.files.get("external_file")
    rubric = state.get("rubric_json")
    if not upload:
        return jsonify({"error": "未收到成绩单"}), 400
    if not rubric:
        return jsonify({"error": "请先生成评分细则"}), 400
    excluded = [c["name"] for c in rubric.get("criteria", []) if c.get("excluded")]
    if not excluded:
        return jsonify({"error": "当前没有停用的维度，无需上传成绩单"}), 400
    try:
        grid = _read_sheet(upload)
    except (ValueError, OSError) as exc:
        return jsonify({"error": str(exc)}), 400
    if not grid:
        return jsonify({"error": "成绩单为空"}), 400

    headers = [str(h or "").strip() for h in grid[0]]
    columns = {}
    for name in excluded:
        index = next((i for i, h in enumerate(headers) if h == name), None)
        if index is None:
            index = next((i for i, h in enumerate(headers)
                          if h and (name in h or h in name)), None)
        if index is not None:
            columns[name] = index
    if not columns:
        return jsonify({"error": f"成绩单里没有可匹配的维度（需要有列名：{'、'.join(excluded)}）"}), 400

    key_col = next((i for i, h in enumerate(headers) if "学号" in h), None)
    name_col = next((i for i, h in enumerate(headers) if "姓名" in h), None)
    if key_col is None and name_col is None:
        return jsonify({"error": "成绩单需要「学号」或「姓名」列"}), 400

    scores = {name: {} for name in columns}
    raw = {name: {} for name in columns}
    for row in grid[1:]:
        key = ""
        for col in (key_col, name_col):
            if col is not None and col < len(row) and str(row[col] or "").strip():
                key = str(row[col]).strip()
                break
        if key.endswith(".0"):
            key = key[:-2]
        if not key:
            continue
        for name, index in columns.items():
            if index >= len(row):
                continue
            value = _score_value(row[index])
            if value is None:
                continue
            scores[name][key] = value
            raw[name][key] = str(row[index])

    if not any(scores.values()):
        return jsonify({"error": "成绩单里没有可用分数（A–E 或 0–100）"}), 400

    data = {"file": sanitize_filename(upload.filename), "at": int(time.time()),
            "scores": scores, "raw": raw}
    external_scores.save(session_dir(), data)
    state["external"] = data
    applied = _apply_external()
    _export_all()
    return jsonify({"ok": True, "file": data["file"], "applied": applied,
                    "dimensions": {name: len(values) for name, values in scores.items()}})


@app.route("/external/clear", methods=["POST"])
def external_clear():
    """清除课外成绩单：删文件、清状态，并把已并入结果的 external 分数回滚、重算小计。"""
    if _busy():
        return jsonify({"error": "有任务正在进行，稍后再清除"}), 409
    if not state.get("external"):
        return jsonify({"ok": True, "cleared": 0})
    storage.remove(external_scores.path_for(session_dir()))
    state["external"] = {}
    grader = state.get("grader")
    rubric = state.get("rubric_json")
    if grader is not None and rubric:
        for result in grader.results:
            grading = result.get("grading") or {}
            scores = grading.get("scores")
            if isinstance(scores, dict) and any((s or {}).get("external") for s in scores.values()):
                grading["scores"] = {n: s for n, s in scores.items() if not (s or {}).get("external")}
                grading["subtotals"] = rubric_io.compute_subtotals(grading["scores"], rubric)
                result["grading"] = grading
        _export_all()
    return jsonify({"ok": True, "cleared": 1})


# ---------------------------------------------------------------- 样本校准
def _pilot_path(base=None):
    return os.path.join(base or session_dir(), "pilot", "pilot.json")


def _load_pilot(base=None):
    data = storage.read_json(_pilot_path(base), {}) or {}
    data.setdefault("reports", [])
    data.setdefault("confirmed_at", 0)
    data.setdefault("rubric_hash", "")
    return data


def _save_pilot(data, base=None):
    storage.write_json(_pilot_path(base), data, indent=2)


def _rubric_hash():
    criteria = (state.get("rubric_json") or {}).get("criteria") or []
    raw = json.dumps(criteria, sort_keys=True, ensure_ascii=False)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12]


def _rubric_confirmed():
    pilot = _load_pilot()
    return bool(pilot.get("confirmed_at")) and pilot.get("rubric_hash") == _rubric_hash()


def _pilot_text_path(entry, base=None):
    return os.path.join(base or session_dir(), entry.get("text") or "")


@app.route("/pilot/upload", methods=["POST"])
def pilot_upload():
    files = request.files.getlist("pilot_files")
    if not files:
        return jsonify({"error": "未收到测试报告"}), 400
    data = _load_pilot()
    if len(data["reports"]) + len(files) > 5:
        return jsonify({"error": "样本校准最多 5 份报告"}), 400
    upload_dir = os.path.join(session_sub("pilot"), "uploads")
    os.makedirs(upload_dir, exist_ok=True)
    extractor = ReportExtractor(os.path.join(session_sub("pilot"), "reports"), state["config"])
    for upload in files:
        name = sanitize_filename(upload.filename)
        path = os.path.join(upload_dir, name)
        upload.save(path)
        result = extractor._process_single(path)
        if result.get("status") != "ok":
            return jsonify({"error": f"{name} 解析失败：{result.get('reason') or '无文本'}"}), 400
        text_path = os.path.join(result["meta"]["output_dir"], "text.md")
        data["reports"] = [r for r in data["reports"] if r.get("filename") != name]
        data["reports"].append({
            "filename": name,
            "chars": len(_read_text(text_path)),
            "text": os.path.relpath(text_path, session_dir()),
            "ai": {}, "teacher": {}, "status": "pending",
        })
    _save_pilot(data)
    return jsonify({"ok": True, "count": len(data["reports"])})


@app.route("/pilot/run", methods=["POST"])
def pilot_run():
    if not state.get("rubric_json"):
        return jsonify({"error": "请先生成评分细则"}), 400
    data = _load_pilot()
    if not data["reports"]:
        return jsonify({"error": "请先上传测试报告"}), 400
    if task_running():
        return jsonify({"error": "有任务正在进行，请稍候"}), 409
    # 请求线程内构造 Grader 与路径：MeteredLLM 在此捕获正确的会话，用量不串课
    session_path = session_dir()
    pilot_dir = os.path.join(session_path, "pilot")
    grader = Grader(state["rubric_json"], llm_client=llm(state["config"]),
                    config=state["config"],
                    output_dir=os.path.join(pilot_dir, "grading"),
                    kb_session_dir=os.path.join(pilot_dir, "kb"))
    t = task_start("pilot", "/app/rubric")

    def work():
        try:
            total = len(data["reports"])
            for index, entry in enumerate(data["reports"]):
                if task_cancelled(t):
                    task_finish(t, "cancelled")
                    return
                task_update(t, progress=90 * index / total, detail=entry["filename"])
                text = _read_text(_pilot_text_path(entry, session_path))
                chunks = _chunks(text)
                assignments = grader.assign_items(
                    [{"id": f"p{i}", "text": chunk[:300]} for i, chunk in enumerate(chunks)],
                    shots_query=text[:1500])
                outcome = grader.grade_single(entry["filename"], text)
                scores = (outcome["result"].get("grading") or {}).get("scores") or {}
                entry["ai"] = {name: detail.get("level") for name, detail in scores.items()
                               if detail.get("level")}
                entry["chunks"] = chunks
                entry["assignments"] = assignments
                entry["status"] = outcome["result"].get("status")
                task_update(t, done=index + 1)
            _save_pilot(data, session_path)
            task_update(t, progress=100, stage="完成",
                        detail=f"已用模型评完 {len(data['reports'])} 份样本")
            task_finish(t, "done")
        except Exception as exc:
            task_finish(t, "cancelled" if task_cancelled(t) else "error", exc)

    threading.Thread(target=work, daemon=True).start()
    return jsonify({"ok": True})


@app.route("/pilot/teacher", methods=["POST"])
def pilot_teacher():
    payload = request.get_json(silent=True) or {}
    levels = payload.get("levels") or {}
    data = _load_pilot()
    for entry in data["reports"]:
        incoming = levels.get(entry["filename"]) or {}
        entry["teacher"] = {name: level for name, level in incoming.items()
                            if level in VALID_LEVELS}
    _save_pilot(data)
    return jsonify({"ok": True})


@app.route("/pilot/delete", methods=["POST"])
def pilot_delete():
    """删除单份校准样本：pilot.json 条目、pilot/reports 提取目录、pilot/uploads 原件。"""
    payload = request.get_json(silent=True) or {}
    filename = str(payload.get("filename", "")).strip()
    data = _load_pilot()
    entry = next((e for e in data["reports"] if e.get("filename") == filename), None)
    if entry is None:
        return jsonify({"error": "样本不存在"}), 404
    data["reports"] = [e for e in data["reports"] if e.get("filename") != filename]
    _save_pilot(data)
    base = session_dir()
    if entry.get("text"):
        storage.remove_tree(os.path.dirname(os.path.join(base, entry["text"])))
    try:
        os.remove(os.path.join(base, "pilot", "uploads", filename))
    except OSError:
        pass
    return jsonify({"ok": True, "deleted": filename})


@app.route("/pilot/text/<path:filename>")
def pilot_text(filename):
    data = _load_pilot()
    entry = next((item for item in data["reports"] if item.get("filename") == filename), None)
    if not entry:
        return jsonify({"error": "样本不存在"}), 404
    return jsonify({"text": _read_text(_pilot_text_path(entry))})


@app.route("/pilot/confirm", methods=["POST"])
def pilot_confirm():
    if not state.get("rubric_json"):
        return jsonify({"error": "请先生成评分细则"}), 400
    data = _load_pilot()
    data["confirmed_at"] = time.time()
    data["rubric_hash"] = _rubric_hash()
    _sync_dimension_collections([c["name"] for c in state["rubric_json"]["criteria"]])
    _save_pilot(data)
    return jsonify({"ok": True, "confirmed": True})


@app.route("/api/pilot")
def api_pilot():
    data = _load_pilot()
    rubric = state.get("rubric_json") or {}
    names = [c["name"] for c in rubric.get("criteria", [])]
    pairs, per_dim = [], {name: [] for name in names}
    for entry in data["reports"]:
        for name in names:
            ai, teacher = (entry.get("ai") or {}).get(name), (entry.get("teacher") or {}).get(name)
            if ai and teacher:
                pairs.append((ai, teacher))
                per_dim.setdefault(name, []).append((ai, teacher))
    metrics = {"cohen_kappa": agreement.cohen_kappa(pairs),
               "weighted_kappa": agreement.weighted_kappa(pairs),
               "krippendorff_alpha": agreement.krippendorff_alpha(pairs),
               "n": len(pairs)}
    confirmed = bool(data.get("confirmed_at")) and data.get("rubric_hash") == _rubric_hash()
    reports = [{key: value for key, value in entry.items() if key != "text"}
               for entry in data["reports"]]
    return jsonify({
        "reports": reports,
        "agreement": metrics,
        "dimensions": [{"name": name,
                        "n": len(per_dim.get(name) or []),
                        "cohen_kappa": agreement.cohen_kappa(per_dim.get(name) or [])}
                       for name in names],
        "confirmed": confirmed,
        "confirmedAt": data.get("confirmed_at") or 0,
        "needsConfirm": bool(data.get("confirmed_at")) and not confirmed,
        "hasRubric": bool(rubric),
    })


@app.route("/task_status")
def task_status():
    t = state.get("task") or {}
    payload = dict(t)
    if t.get("started"):
        payload["elapsed"] = round(time.time() - t["started"], 1)
    return jsonify(payload)


@app.route("/task_cancel", methods=["POST"])
def task_cancel():
    t = state.get("task")
    if t and t["status"] == "running":
        t["cancel"] = True
        t["stage"] = "正在取消（等待当前步骤结束）…"
    return jsonify({"ok": True})



@app.route("/rubric_gen_cell", methods=["POST"])
def rubric_gen_cell():
    data = request.get_json(silent=True) or {}
    name = str(data.get("name", "")).strip()
    fields = [f for f in (data.get("fields") or [])
              if f in {"weight", "doc_ref", "A", "B", "C", "D", "E"}]
    if not name or not fields:
        return jsonify({"error": "需要维度名称和至少一个待补全字段"}), 400
    try:
        filled = RubricGenerator(llm_client=llm(state["config"])).generate_cells(
            state.get("criteria_doc", ""), name, data.get("context") or {}, fields)
    except Exception as e:
        return jsonify({"error": f"生成失败：{e}"}), 422
    return jsonify(filled)


@app.route("/rubric_table_save", methods=["POST"])
def rubric_table_save():
    if _busy():
        return jsonify({"error": "评分进行中，暂不能修改细则"}), 409
    data = request.get_json(silent=True) or {}
    criteria = [c for c in (data.get("criteria") or []) if not c.get("_deleted")]
    if not criteria:
        return jsonify({"error": "至少保留一个评分维度"}), 400
    if any(not str(c.get("name", "")).strip() for c in criteria):
        return jsonify({"error": "存在未命名维度，请填写维度名称"}), 400
    if not data.get("confirm_reset") and _has_downstream():
        return jsonify({"needReset": True,
                        "error": "保存细则会清除已提取的报告、评分结果与向量库（上传原件保留），请确认后重试"}), 409
    warnings = rubric_io.check_vague_and_refs(criteria)
    if warnings and not data.get("force"):
        return jsonify({"warnings": warnings}), 409
    previous = {c["name"]: c for c in (state.get("rubric_json") or {}).get("criteria", [])}
    for criterion in criteria:
        old = previous.get(criterion.get("name"))
        if not old:
            continue
        changed = any(json.dumps(criterion.get(key), sort_keys=True, ensure_ascii=False)
                      != json.dumps(old.get(key), sort_keys=True, ensure_ascii=False)
                      for key in ("levels", "doc_ref", "mechanical", "excluded"))
        if changed:
            criterion["edited"] = True
    rubric_json = {"criteria": rubric_io.normalize_weights(criteria)}
    state["rubric_json"] = rubric_json
    state["rubric"] = rubric_io.criteria_to_markdown(rubric_json)
    state["rubric_warnings"] = []
    _save_rubric(rubric_json)  # 内部重置下游并作废校准样本模型判级
    return jsonify({"ok": True})


# ---------------------------------------------------------------- 校准模式
@app.route("/calibrate_upload", methods=["POST"])
def calibrate_upload():
    report_file = request.files.get("calibration_report")
    if not report_file:
        return jsonify({"error": "未收到校准报告"}), 400
    if not state.get("rubric_json"):
        return redirect("/app/rubric")
    filepath = os.path.join(session_sub("uploads"),
                            "calibration_" + sanitize_filename(report_file.filename))
    report_file.save(filepath)
    extraction = ReportExtractor(session_sub("calibration"), state["config"])._process_single(filepath)
    if extraction.get("status") != "ok":
        return jsonify({"error": f"校准报告解析失败：{extraction.get('reason')}"}), 400
    text = _read_text(os.path.join(extraction["meta"]["output_dir"], "text.md"))
    chunks = _chunks(text)
    grader = Grader(state["rubric_json"], llm_client=llm(state["config"]), config=state["config"],
                    similarity_scope=state.get("session"),
                    kb_session_dir=os.path.join(session_dir(), "calibration_kb"))
    assignments = grader.assign_items([{"id": f"p{i}", "text": c[:300]}
                                       for i, c in enumerate(chunks)], shots_query=text[:1500])
    state["calibration"] = {"text": text, "chunks": chunks, "assignments": assignments,
                            "filename": extraction["meta"].get("filename", report_file.filename)}
    storage.write_json(os.path.join(session_dir(), "calibration_workspace.json"),
                       state["calibration"])
    return redirect("/app/calibration")




# ---------------------------------------------------------------- 报告
@app.route("/upload_reports", methods=["POST"])
def upload_reports():
    files = request.files.getlist("report_files")
    if not files:
        return jsonify({"error": "未收到文件"}), 400
    if _busy():
        return jsonify({"error": "有任务正在进行，请稍候或先取消"}), 409
    upload_dir = session_sub("uploads")
    saved = []
    for f in files:
        name = sanitize_filename(f.filename)
        f.save(os.path.join(upload_dir, name))
        saved.append(name)
    state["report_uploads"] = saved
    return jsonify({"ok": True, "count": len(saved), "filenames": saved})


@app.route("/process_reports", methods=["POST"])
def process_reports():
    if (state.get("grading_status") or {}).get("phase") == "running":
        return redirect("/app/results")
    if task_running():
        return redirect("/app/workbench")
    # 请求线程内捕获会话目录、已提取列表与去重基准；线程只写捕获值
    sid = state["session"]
    reports_dir = session_sub("reports")
    upload_dir = session_sub("uploads")
    config = state["config"]
    fields = _student_fields()
    report_uploads = list(state.get("report_uploads") or [])
    prev = [r for r in state["extracted_reports"] if r.get("status") == "ok"]
    student_llm = llm(config)
    t = task_start("reports", "/app/reports")

    def work():
        try:
            prev_names = {r["meta"]["filename"] for r in prev}
            prev_students = {r["meta"]["filename"]: r["meta"].get("student")
                             for r in prev if r["meta"].get("student")}
            seen_hashes = {h for h in (_text_hash(r) for r in prev) if h}
            extractor = ReportExtractor(reports_dir, config)
            queued = [os.path.join(upload_dir, n) for n in report_uploads
                      if os.path.isfile(os.path.join(upload_dir, n))]
            if queued:
                all_files = queued
            else:
                all_files = [p for p in collect_report_files(upload_dir)
                             if not os.path.basename(p).startswith(
                                 ("calibration_", "criteria_", "experiment_"))]
            total = len(all_files)
            if not total:
                task_finish(t, "error", "uploads 目录中没有可处理的报告文件")
                return
            task_update(t, total=total, stage="提取报告内容")
            kept, dup = [], 0
            for i, fp in enumerate(all_files):
                if task_cancelled(t):
                    task_finish(t, "cancelled")
                    return
                name = os.path.basename(fp)
                task_update(t, progress=90 * i / max(total, 1), detail=name)
                entry = extractor._process_single(fp)
                if entry.get("status") == "ok":
                    fn = entry["meta"]["filename"]
                    if fn in prev_students:
                        entry["meta"]["student"] = prev_students[fn]
                    elif not entry["meta"].get("student"):
                        task_update(t, stage="识别学生信息", detail=name)
                        entry["meta"]["student"] = _extract_student(entry, fields, student_llm)
                        _save_meta(entry)
                    digest = _text_hash(entry)
                    if digest and fn not in prev_names and digest in seen_hashes:
                        _discard_report(entry, upload_dir)
                        dup += 1
                        task_update(t, done=i + 1)
                        continue
                    if digest:
                        seen_hashes.add(digest)
                kept.append(entry)
                task_update(t, done=i + 1)
            if task_cancelled(t):
                task_finish(t, "cancelled")
                return
            # 已切走会话时不动新会话内存：提取物已落旧会话目录，切回时 hydrate 自磁盘重建
            if state.get("session") == sid:
                state["extracted_reports"] = kept
                state["report_uploads"] = []
            detail = f"共 {len(kept)} 份" + (f"，自动去除 {dup} 份重复" if dup else "")
            task_update(t, progress=100, stage="完成", detail=detail)
            task_finish(t, "done")
        except Exception as e:
            task_finish(t, "cancelled" if task_cancelled(t) else "error", e)

    threading.Thread(target=work, daemon=True).start()
    return redirect("/app/workbench")


def _save_meta(entry):
    storage.write_json(os.path.join(entry["meta"]["output_dir"], "meta.json"),
                       entry["meta"], indent=2)


@app.route("/reports/delete/<path:filename>", methods=["POST"])
def reports_delete(filename):
    """删除单份报告及其全部衍生物：提取目录、上传原件、评分结果、向量记录、终审/重评记录。"""
    if _busy():
        return jsonify({"error": "有任务正在进行，稍后再删除"}), 409
    entry = _find_entry(filename)
    if entry is None:
        return jsonify({"error": "报告不存在"}), 404
    grader = state.get("grader")
    if grader is not None:
        grader.results = [r for r in grader.results if r["report_id"] != filename]
        grader.export_results()
        grader._purge_report_vectors(filename)
    state["approved_ids"] = _approved_ids() - {filename}
    save_approved(state["approved_ids"])
    if state["corrections"].pop(filename, None) is not None:
        storage.write_json(os.path.join(session_dir(), "corrections.json"), state["corrections"])
    storage.remove_tree(entry["meta"]["output_dir"])
    origin = _origin_path(filename)
    if origin:
        try:
            os.remove(origin)
        except OSError:
            pass
    state["extracted_reports"] = [r for r in state["extracted_reports"]
                                  if r["meta"]["filename"] != filename]
    state["report_uploads"] = [n for n in (state.get("report_uploads") or []) if n != filename]
    if grader is not None:
        _export_all()
    return jsonify({"ok": True, "deleted": filename})


def _doc_author(entry):
    """读取原文件的文档属性作者/标题（PDF 元数据 / docx 核心属性）。"""
    src = entry.get("file", "")
    try:
        if os.path.splitext(src)[1].lower() == ".pdf" and os.path.exists(src):
            import pymupdf
            doc = pymupdf.open(src)
            meta = doc.metadata or {}
            doc.close()
            return str(meta.get("author") or "").strip(), str(meta.get("title") or "").strip()
        if os.path.splitext(src)[1].lower() == ".docx" and os.path.exists(src):
            from docx import Document
            props = Document(src).core_properties
            return str(props.author or "").strip(), str(props.title or "").strip()
    except Exception:
        pass
    return "", ""


def _extract_student(entry, fields, client=None):
    head = _read_text(os.path.join(entry["meta"]["output_dir"], "text.md"))[:2500]
    author, title = _doc_author(entry)
    prompt = ("从下面这份实验报告的文件名、文档属性和开头内容中提取学生信息。报告可能是英文的："
              "作者行常形如「Authors: Ryan D. Martin (Queens' University)」或「姓名：张三　学号：20230001」。"
              f"字段列表：{fields}。找不到的字段填空字符串。"
              "若列出多位作者，取第一位填入姓名字段。学号/班级/学院常见于标题下方、页眉或作者行附近。"
              "只输出一个 JSON 对象，不要输出其他文字。\n\n"
              f"文件名：{entry['meta']['filename']}\n文档属性作者：{author or '（无）'}\n"
              f"文档属性标题：{title or '（无）'}\n\n报告开头内容：\n{head}")
    try:
        data = _json_or_empty((client or llm(state["config"]))
                              .chat([{"role": "user", "content": prompt}]))
    except Exception:
        data = {}
    return {k: str(data.get(k, "") or "").strip() for k in fields}


def _json_or_empty(text):
    m = re.search(r"\{.*\}", text or "", re.S)
    return json.loads(m.group(0)) if m else {}


def _text_hash(entry):
    norm = re.sub(r"\s+", "", _report_text(entry))
    return hashlib.sha1(norm.encode("utf-8")).hexdigest() if norm else None


def _discard_report(entry, upload_dir):
    import shutil
    shutil.rmtree(entry["meta"]["output_dir"], ignore_errors=True)
    src = os.path.abspath(entry.get("file", ""))
    roots = (os.path.abspath(upload_dir), os.path.abspath(session_sub("reports")))
    if src and os.path.exists(src) and any(src.startswith(r) for r in roots):
        try:
            os.remove(src)
        except OSError:
            pass



@app.route("/report_text/<path:filename>")
def report_text(filename):
    entry = _find_entry(filename)
    if not entry:
        return jsonify({"error": "报告不存在"}), 404
    chunks = _chunks(_report_text(entry), 10 ** 6)
    pages = [chunks[i:i + 4] for i in range(0, len(chunks), 4)] or [["（此报告无文本内容）"]]
    return jsonify({"filename": filename, "pages": pages})


@app.route("/report_original/<path:filename>")
def report_original(filename):
    path = _origin_path(filename)
    if not path:
        return "原文文件不存在", 404
    return send_file(path)


def _build_dim_blocks(rubric_json, text, result, fallback_images=()):
    result = result or {}
    chunks = result.get("_report_chunks") or _chunks(text)
    assignments = result.get("text_assignments") or {}
    if result.get("_report_chunks") and chunks != _chunks(text):
        # 旧数据按被公式拆散的块保存过：重新分块并把挂载关系按文本映射到新块
        fresh = _chunks(text)
        remapped = {}
        for chunk_id, assignment in assignments.items():
            index = int(chunk_id[1:]) if chunk_id[1:].isdigit() else -1
            if not 0 <= index < len(chunks):
                continue
            moved = next((j for j, chunk in enumerate(fresh) if chunks[index].strip() in chunk), None)
            if moved is not None:
                remapped[f"p{moved}"] = assignment
        chunks, assignments = fresh, remapped
    img_assign = result.get("image_assignments") or {}
    images = result.get("images") or list(fallback_images)
    scores = (result.get("grading") or {}).get("scores", {})
    blocks, used_chunks, used_imgs = [], set(), set()
    for c in rubric_json.get("criteria", []):
        name = c["name"]
        dim_chunks = [chunks[int(aid[1:])] for aid, a in assignments.items()
                      if (a or {}).get("criterion") == name and aid[1:].isdigit()
                      and int(aid[1:]) < len(chunks)]
        used_chunks.update(aid for aid, a in assignments.items() if (a or {}).get("criterion") == name)
        dim_imgs = [im for im in images if (img_assign.get(im["filename"]) or {}).get("criterion") == name]
        used_imgs.update(im["filename"] for im in dim_imgs)
        blocks.append({"criterion": name, "weight": c.get("weight"), "mechanical": c.get("mechanical"),
                       "doc_ref": c.get("doc_ref"), "chunks": dim_chunks, "images": dim_imgs,
                       "score": scores.get(name)})
    rest_chunks = [chunks[int(aid[1:])] for aid in
                   [f"p{i}" for i in range(len(chunks))] if aid not in used_chunks]
    rest_images = [im for im in images if im["filename"] not in used_imgs]
    if rest_chunks or rest_images:
        blocks.append({"criterion": "未挂载内容", "weight": None, "mechanical": False, "doc_ref": "",
                       "chunks": rest_chunks, "images": rest_images, "score": None})
    return blocks


@app.route("/report/<path:filename>")
def report_detail(filename):
    entry = _find_entry(filename)
    if not entry or entry.get("status") != "ok":
        return "报告不存在", 404
    return redirect(f"/app/report/{quote(filename, safe='/')}")


@app.route("/report_image/<path:filename>/<path:imgname>")
def report_image(filename, imgname):
    entry = _find_entry(filename)
    if not entry:
        return "未找到", 404
    path = os.path.join(entry["meta"]["output_dir"], "images", os.path.basename(imgname))
    if not os.path.isfile(path):
        return "图片不存在", 404
    return send_file(path)


@app.route("/update_student/<path:filename>", methods=["POST"])
def update_student(filename):
    entry = _find_entry(filename)
    if not entry:
        return jsonify({"error": "报告不存在"}), 404
    student = entry["meta"].get("student") or {}
    for field in _student_fields():
        if f"student::{field}" in request.form:
            student[field] = request.form.get(f"student::{field}", "").strip()
    entry["meta"]["student"] = student
    _save_meta(entry)
    _export_all()
    return redirect(request.form.get("next") or "/app/reports")


# ---------------------------------------------------------------- 学生信息与花名册
def _roster_key(roster):
    for header in roster["headers"]:
        if "学号" in header:
            return header
    for header in roster["headers"]:
        if "姓名" in header:
            return header
    return roster["headers"][0] if roster["headers"] else None


def _roster_index(roster):
    key = _roster_key(roster)
    values = {}
    for row in roster["rows"]:
        value = str(row.get(key, "") or "").strip()
        if value:
            values.setdefault(value, row)
    return key, values


def _classify_students(reports=None, roster=None):
    """返回 (students, suspicious, missing)。不传参数时读当前会话。"""
    reports = state["extracted_reports"] if reports is None else reports
    students = {r["meta"]["filename"]: (r["meta"].get("student") or {})
                for r in reports if r.get("status") == "ok"}
    roster = load_roster() if roster is None else roster
    if not roster:
        return students, [], []
    key, values = _roster_index(roster)
    name_key = next((h for h in roster["headers"] if "姓名" in h and h != key), None)
    suspicious, matched = [], []
    for filename, student in students.items():
        value = str(student.get(key, "") or "").strip()
        if not value:
            suspicious.append({"report_id": filename, "student": student, "reason": "未提取到学生信息"})
        elif value not in values:
            suspicious.append({"report_id": filename, "student": student,
                               "reason": f"{key}「{value}」不在花名册中"})
        else:
            if name_key:
                roster_name = str(values[value].get(name_key, "") or "").strip()
                report_name = str(student.get(name_key, "") or "").strip()
                if roster_name and report_name and roster_name != report_name:
                    suspicious.append({"report_id": filename, "student": student,
                                       "reason": f"{key}匹配但{name_key}不一致（名册：{roster_name}）"})
            matched.append(value)
    missing = [row for row in roster["rows"] if str(row.get(key, "") or "").strip() not in matched]
    return students, suspicious, missing


def _export_all(grader=None, reports=None, roster=None, external=None, rubric=None):
    """导出成绩 Excel 与被标记名单。无参时读当前会话（请求线程）；
    评分线程必须传捕获的参数，避免评分完成后写错会话。"""
    if grader is None:
        grader = state["grader"]
        reports = state["extracted_reports"]
        roster = load_roster()
        external = state.get("external") or {}
        rubric = state.get("rubric_json")
    if not grader:
        return
    if external and rubric:
        students = {r["meta"]["filename"]: (r["meta"].get("student") or {})
                    for r in reports or []}
        external_scores.apply_to_results(grader.results, rubric, external, students.get)
    students_map, suspicious, missing = _classify_students(reports=reports, roster=roster)
    grader.export_results()
    grader.export_excel(roster=roster, students=students_map,
                        suspicious=suspicious, missing=missing)
    _export_flagged(suspicious, roster=roster, grading_dir=grader.output_dir)


def _export_flagged(suspicious, roster=None, grading_dir=None):
    import openpyxl
    cols = _student_fields() if roster is None else roster["headers"]
    grading_dir = grading_dir or session_sub("grading")
    path = os.path.join(grading_dir, "flagged_students.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "被标记人员"
    ws.append(cols + ["报告ID", "标记原因"])
    for s in suspicious:
        ws.append([(s["student"] or {}).get(c, "") for c in cols] + [s["report_id"], s["reason"]])
    wb.save(path)



@app.route("/roster/upload", methods=["POST"])
def roster_upload():
    f = request.files.get("roster_file")
    if not f:
        return jsonify({"error": "未收到文件"}), 400
    ext = os.path.splitext(f.filename)[1].lower()
    if ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), read_only=True)
        grid = [list(row) for row in wb.active.iter_rows(values_only=True)]
        wb.close()
    elif ext == ".csv":
        import csv
        grid = list(csv.reader(io.StringIO(f.read().decode("utf-8-sig", errors="replace"))))
    else:
        return jsonify({"error": "仅支持 .xlsx / .csv"}), 400
    if not grid:
        return jsonify({"error": "文件为空"}), 400
    headers = [str(h or "").strip() or f"列{i + 1}" for i, h in enumerate(grid[0])]
    rows = [{headers[i]: ("" if c[i] is None else str(c[i]).strip())
             for i in range(min(len(headers), len(c)))}
            for c in grid[1:] if any(x not in (None, "") for x in c)]
    save_roster({"headers": headers, "rows": rows})
    return redirect("/app/roster")


@app.route("/roster/save", methods=["POST"])
def roster_save():
    try:
        data = json.loads(request.form.get("data", ""))
        headers = [str(h).strip() for h in data.get("headers", []) if str(h).strip()]
        if not headers:
            return jsonify({"error": "表头不能为空"}), 400
        rows = [{h: str(row.get(h, "") or "").strip() for h in headers}
                for row in data.get("rows", [])]
    except (ValueError, AttributeError):
        return jsonify({"error": "数据格式错误"}), 400
    save_roster({"headers": headers, "rows": rows})
    return redirect("/app/roster")


# ---------------------------------------------------------------- MCP 插件
def _mcp_references(query, mount=None, sid=None):
    """拉学科参考资料。mount/sid 由后台线程传入捕获值；只在仍是当前会话时写 mcp_status。"""
    mount = mount if mount is not None else (state.get("mcp") or {})
    if not mount.get("servers"):
        return ""
    servers = mcp_store.load_servers()
    text, status = mcp_client.collect(servers, mount, query)
    if sid is None or sid == state.get("session"):
        state["mcp_status"] = status
    return text



@app.route("/mcp/add", methods=["POST"])
def mcp_add():
    try:
        cfg = mcp_store.parse_form(request.form)
    except ValueError as e:
        return redirect("/app/mcp")
    mcp_store.upsert(request.form.get("name", "").strip(), cfg)
    return redirect("/app/mcp")


@app.route("/mcp/remove", methods=["POST"])
def mcp_remove():
    name = request.form.get("name", "").strip()
    mcp_store.remove(name)
    mount = state.get("mcp") or {}
    mount["servers"] = [s for s in mount.get("servers") or [] if s != name]
    save_mcp_mount(mount)
    return redirect("/app/mcp")


@app.route("/mcp/mount", methods=["POST"])
def mcp_mount():
    mount = {
        "servers": request.form.getlist("servers"),
        "resources": [p.strip() for p in (request.form.get("resources") or "").splitlines() if p.strip()],
        "search_tool": (request.form.get("search_tool") or "").strip(),
        "prompt": (request.form.get("prompt") or "").strip(),
        "max_chars": max(200, min(20000, int(request.form.get("max_chars") or 4000))),
    }
    save_mcp_mount(mount)
    return redirect("/app/mcp")


@app.route("/mcp/test/<name>")
def mcp_test(name):
    cfg = mcp_store.load_servers().get(name)
    if not cfg:
        return jsonify({"error": "插件不存在"}), 404
    try:
        return jsonify(mcp_client.describe(cfg))
    except Exception as e:
        return jsonify({"error": str(e)[:300]}), 200


@app.route("/mcp/refresh", methods=["POST"])
def mcp_refresh():
    query = " ".join(c["name"] for c in (state.get("rubric_json") or {}).get("criteria", [])) or "实验报告"
    state["mcp_refs"] = _mcp_references(query)
    return redirect("/app/mcp")


# ---------------------------------------------------------------- 评分
def _image_entries(entry):
    """图片 → 归维度用的条目。默认用图片上下文；开启 grading.ocr_images 才 OCR（结果缓存）。"""
    from .utils.ocr_engine import OCREngine
    enabled = bool((state["config"] or {}).get("grading", {}).get("ocr_images", False))
    report_dir = entry["meta"]["output_dir"]
    cache_dir = os.path.join(report_dir, ".ocr")
    os.makedirs(cache_dir, exist_ok=True)
    engine = None
    entries = []
    for image in entry.get("images") or []:
        name = image["filename"]
        cache_path = os.path.join(cache_dir, name + ".txt")
        if not enabled:
            entries.append({"id": name, "text": image.get("context") or name})
            continue
        if os.path.exists(cache_path):
            text = _read_text(cache_path)
        else:
            text = ""
            if engine is None or engine is False:
                try:
                    engine = OCREngine.get((state["config"] or {}).get("ocr", {}))
                except Exception:
                    engine = False
            if engine:
                try:
                    text = engine.recognize_file(os.path.join(report_dir, "images", name))
                except Exception:
                    text = ""
            storage.write_text(cache_path, text)
        entries.append({"id": name, "text": text or image.get("context") or name})
    return entries


def _image_descriptions(entries):
    return "\n".join(f"- 图 {e['id']}：{e['text'][:200]}" for e in entries[:20]) or ""


def _grade_one(grader, entry, references="", corrections=None):
    from .web import usage as usage_store

    text = _report_text(entry)
    if not text.strip():
        return None
    # corrections 由评分轮次捕获传入（会话隔离）；同步调用缺省读当前会话
    corrections = corrections if corrections is not None else state["corrections"]
    filename = entry["meta"]["filename"]
    # 这段期间的所有模型调用都记到这份报告名下，用于「每份报告成本」
    usage_store.set_report(filename)
    try:
        images = entry.get("images") or []
        image_entries = _image_entries(entry)
        chunks = _chunks(text)
        chunk_entries = [{"id": f"p{i}", "text": c[:300]} for i, c in enumerate(chunks)]
        correction_text = corrections.get(filename)
        grade_text = text + (f"\n\n[教师修改意见，本次重评须重点回应] {correction_text}"
                             if correction_text else "")

        # 重评且正文、维度名都没变：归类结果可复用，省掉一次模型调用
        criterion_names = {c["name"] for c in grader.rubric.get("criteria", [])}
        previous = next((r for r in grader.results if r.get("report_id") == filename), None)
        reusable = False
        if previous and previous.get("_report_text") == text:
            prior_text = previous.get("text_assignments") or {}
            prior_images = previous.get("image_assignments") or {}
            chunk_ids = {f"p{i}" for i in range(len(chunks))}
            image_ids = {e["id"] for e in image_entries}
            reusable = (set(prior_text) == chunk_ids and set(prior_images) == image_ids
                        and all((a or {}).get("criterion") in criterion_names
                                for a in list(prior_text.values()) + list(prior_images.values())))

        def run_assign():
            usage_store.set_report(filename)
            try:
                return grader.assign_items(image_entries + chunk_entries, shots_query=text[:1500])
            finally:
                usage_store.set_report(None)

        def run_grade():
            usage_store.set_report(filename)
            try:
                return grader.grade_single(
                    filename, grade_text, image_descriptions=_image_descriptions(image_entries),
                    references=references)
            finally:
                usage_store.set_report(None)

        # 归类与判级互不依赖：并行可省掉一次模型往返；单路模式保持串行，避免本地模型被并发拖慢
        if reusable:
            assignments = {**previous.get("text_assignments", {}),
                           **previous.get("image_assignments", {})}
            outcome = run_grade()
        elif _grading_workers() > 1:
            with ThreadPoolExecutor(max_workers=2, thread_name_prefix="report") as pool:
                assign_future = pool.submit(run_assign)
                grade_future = pool.submit(run_grade)
                assignments = assign_future.result()
                outcome = grade_future.result()
        else:
            assignments = run_assign()
            outcome = run_grade()
        text_assignments = {k: v for k, v in assignments.items() if k.startswith("p")}
        image_assignments = {k: v for k, v in assignments.items() if not k.startswith("p")}
        result = outcome["result"]
        if result.get("status") == "graded":
            # 成功才消费教师意见；llm_failed 时保留，重评继续生效
            corrections.pop(filename, None)
            result["_report_text"] = text
            result["_report_chunks"] = chunks
            result["images"] = images
            result["image_assignments"] = image_assignments
            result["text_assignments"] = text_assignments
        return outcome
    finally:
        usage_store.set_report(None)


def _grading_workers():
    """评分并发数：环境变量 ERG_GRADING_WORKERS 显式覆盖，其次 config 的 grading.workers
    （网页可配），默认 3；一律 clamp 到 1..8。"""
    def clamp(value):
        try:
            return max(1, min(8, int(value)))
        except (TypeError, ValueError):
            return None

    workers = clamp(os.environ.get("ERG_GRADING_WORKERS"))
    if workers:
        return workers
    return clamp((state.get("config") or {}).get("grading", {}).get("workers")) or 3


def _grade_targets(grader, targets, references, status, corrections=None):
    """按报告并行；并发数受 ERG_GRADING_WORKERS 限制，取消后不再派新任务。"""
    workers = _grading_workers()
    if workers <= 1 or len(targets) <= 1:
        for entry in targets:
            if status.get("cancel"):
                return
            status["current"] = entry["meta"]["filename"]
            try:
                _grade_one(grader, entry, references, corrections)
            except Exception as e:  # 单份失败不拖垮整轮
                grader.record_failure(entry["meta"]["filename"], e)
            status["done"] = status.get("done", 0) + 1
        return

    pending, queue = {}, iter(targets)

    def submit(pool):
        if status.get("cancel"):
            return
        try:
            entry = next(queue)
        except StopIteration:
            return
        pending[pool.submit(_grade_one, grader, entry, references, corrections)] = entry

    with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="grade") as pool:
        for _ in range(min(workers, len(targets))):
            submit(pool)
        while pending:
            finished, _ = wait(list(pending), return_when=FIRST_COMPLETED)
            for future in finished:
                entry = pending.pop(future)
                try:
                    future.result()
                except Exception as e:
                    grader.record_failure(entry["meta"]["filename"], e)
                status["done"] = status.get("done", 0) + 1
                status["current"] = entry["meta"]["filename"]
                submit(pool)


def _spawn_grading(sid, grader, targets, corrections, reports, roster,
                  external, rubric_json, mount):
    """会话隔离的评分线程：所有依赖在请求线程内捕获。线程只更新自己持有的
    status dict（登记在 RUNNING_STATUS[sid]），绝不读写 state 里的其他会话数据。"""
    status = {"phase": "running", "done": 0, "total": len(targets), "current": "",
              "cancel": False, "message": "", "workers": _grading_workers()}
    RUNNING_STATUS[sid] = status
    if state.get("session") == sid:
        state["grading_status"] = status
    memory_dir = os.path.join(session_dir(sid), "memory")
    storage.write_json(os.path.join(memory_dir, "grading_status.json"), status)

    def work():
        try:
            references = _mcp_references(" ".join(
                c["name"] for c in (rubric_json or {}).get("criteria", [])),
                mount=mount, sid=sid)
            if state.get("session") == sid:
                state["mcp_refs"] = references
            _grade_targets(grader, targets, references, status, corrections)
            grader.export_results()
            _export_all(grader=grader, reports=reports, roster=roster,
                        external=external, rubric=rubric_json)
            if status.get("cancel"):
                status.update(phase="cancelled", current="",
                              message="已取消，未评分的报告可重新开始评分")
            else:
                status.update(phase="done", current="", message="全部完成")
        except Exception as e:
            status.update({"phase": "error", "message": str(e)[:300], "done": 0, "total": 0})
        finally:
            # 终态与教师意见落存储：重启/切课后仍可看到本轮结果与未消费的意见
            storage.write_json(os.path.join(memory_dir, "grading_status.json"), status)
            storage.write_json(os.path.join(session_dir(sid), "corrections.json"), corrections)

    threading.Thread(target=work, daemon=True).start()


def _student_group(entry, roster):
    """返回 (分组键, 分组标签)。键为空表示未识别学生（每份独立成组）。"""
    student = entry["meta"].get("student") or {}
    if roster:
        key = _roster_key(roster)
        value = str(student.get(key, "") or "").strip()
        if value:
            name_key = next((h for h in roster["headers"] if "姓名" in h and h != key), None)
            name = str(student.get(name_key, "") or "").strip() if name_key else ""
            return value, f"{name}（{value}）" if name else value
    sid = str(student.get("学号", "") or "").strip()
    name = str(student.get("姓名", "") or "").strip()
    if sid:
        return "sid:" + sid, f"{name}（{sid}）" if name else sid
    if name:
        return "name:" + name, name
    return "", "未识别学生"


def _file_size(entry):
    try:
        return os.path.getsize(entry.get("file", "")) if entry.get("file") else 0
    except OSError:
        return 0


def _select_context(error=None):
    graded = _graded_ids()
    approved = _approved_ids()
    candidates = [r for r in state["extracted_reports"]
                  if r.get("status") == "ok" and r["meta"]["filename"] not in approved]
    roster = load_roster()
    groups, order = {}, []
    for r in candidates:
        key, label = _student_group(r, roster)
        gid = key or "unknown"
        if gid not in groups:
            groups[gid] = {"label": label if key else "未识别学生", "multi": bool(key), "reports": []}
            order.append(gid)
        groups[gid]["reports"].append(r)
    for group in groups.values():
        group["reports"].sort(key=lambda r: (r["meta"].get("extract_time", ""), _file_size(r)),
                              reverse=True)
    hint = None
    if roster:
        n = len(roster["rows"])
        hint = f"名册 {n} 人 · 已收报告 {len(candidates)} 份"
        if n != len(candidates):
            hint += f" · 差 {abs(n - len(candidates))} 人，详见「花名册」页"
    statuses = {r["report_id"]: r for r in (state["grader"].results if state["grader"] else [])}
    stale_ids = {rid for rid, r in statuses.items() if _result_stale(r)}
    return {"reports": candidates, "graded_ids": graded, "statuses": statuses,
            "stale_ids": stale_ids, "groups": [groups[g] for g in order],
            "hint": hint, "error": error,
            "has_rubric": bool(state.get("rubric_json") or state.get("rubric")),
            "rubric_confirmed": _rubric_confirmed()}



@app.route("/start_grading", methods=["POST"])
def start_grading():
    wants_json = "application/json" in (request.headers.get("Accept") or "")
    if not state.get("rubric_json") and not state.get("rubric"):
        return jsonify({"error": "尚未配置评分细则"}), 400
    pilot = _load_pilot()
    if state.get("rubric_json") and (not pilot.get("confirmed_at")
                                     or pilot.get("rubric_hash") != _rubric_hash()):
        message = "请先在「评分细则」完成样本校准并确认，再开始评分"
        return jsonify({"error": message, "needConfirm": True}), 409
    with state["grading_lock"]:
        if (state.get("grading_status") or {}).get("phase") == "running":
            if wants_json:
                return jsonify({"ok": True, "started": 0, "running": True,
                                "message": "已有评分任务在跑，正在查看进度"})
            return redirect("/app/results")
        selected = set(request.form.getlist("report_ids"))
        approved = _approved_ids()
        targets = [r for r in state["extracted_reports"]
                   if r.get("status") == "ok" and r["meta"]["filename"] in selected
                   and r["meta"]["filename"] not in approved]
        if not targets:
            if wants_json:
                return jsonify({"ok": False, "started": 0,
                                "message": "没有可评的报告（可能已终审或已被移除）"})
            return redirect("/app/select")
        roster = load_roster()
        seen, conflict = {}, None
        for entry in targets:
            key, label = _student_group(entry, roster)
            if key and key in seen:
                conflict = (f"同一学生 {label} 被勾选了多份：{seen[key]} 和 "
                            f"{entry['meta']['filename']}，每人只能评一份，请返回重新勾选")
                break
            if key:
                seen[key] = entry["meta"]["filename"]
        if conflict:
            if wants_json:
                return jsonify({"ok": False, "started": 0, "message": conflict})
            return redirect("/app/select")
        prev_results = state["grader"].results if state["grader"] else []
        state["grader"] = Grader(state["rubric_json"] or state["rubric"],
                                 llm_client=llm(state["config"]), config=state["config"],
                                 output_dir=session_sub("grading"),
                                 similarity_scope=state.get("session"), kb_session_dir=session_dir())
        # 保留旧结果，评完由 _finish 覆盖；中途取消不至于丢掉上一轮成绩
        state["grader"].results = prev_results
        _spawn_grading(
            sid=state["session"], grader=state["grader"], targets=targets,
            corrections=state["corrections"], reports=state["extracted_reports"],
            roster=roster, external=state.get("external") or {},
            rubric_json=state.get("rubric_json"), mount=dict(state.get("mcp") or {}))
    if wants_json:
        return jsonify({"ok": True, "started": len(targets),
                        "files": [t["meta"]["filename"] for t in targets]})
    return redirect("/app/results")



@app.route("/grading_status")
def grading_status():
    return jsonify(state.get("grading_status") or {"phase": "idle"})


@app.route("/grading_cancel", methods=["POST"])
def grading_cancel():
    status = state.get("grading_status")
    if status and status.get("phase") == "running":
        status["cancel"] = True
    return jsonify({"ok": True})


@app.route("/regrade/<path:filename>", methods=["POST"])
def regrade(filename):
    if not state["grader"]:
        return jsonify({"error": "尚未评分"}), 400
    with state["grading_lock"]:
        if (state.get("grading_status") or {}).get("phase") == "running":
            return redirect("/app/results")
        entry = _find_entry(filename)
        if not entry:
            return jsonify({"error": "报告不存在"}), 404
        corrections = request.form.get("corrections", "").strip()
        if corrections:
            state["corrections"][filename] = corrections
            storage.write_json(os.path.join(session_dir(), "corrections.json"),
                               state["corrections"])
        state["grader"].rubric = _coerce_rubric(state.get("rubric_json") or state.get("rubric") or "")
        state["grader"].results = [r for r in state["grader"].results if r["report_id"] != filename]
        state["approved_ids"] = _approved_ids() - {filename}
        save_approved(state["approved_ids"])
        _spawn_grading(
            sid=state["session"], grader=state["grader"], targets=[entry],
            corrections=state["corrections"], reports=state["extracted_reports"],
            roster=load_roster(), external=state.get("external") or {},
            rubric_json=state.get("rubric_json"), mount=dict(state.get("mcp") or {}))
    return redirect("/app/results")


def _apply_teacher_levels(result, form, level_prefix, reason_prefix=""):
    if result.get("status") != "graded" or not result.get("grading"):
        return False
    scores = result["grading"].get("scores", {})
    changed = False
    for criterion, detail in scores.items():
        if not isinstance(detail, dict):
            continue
        level = form.get(f"{level_prefix}{criterion}")
        if level in VALID_LEVELS and level != detail.get("level"):
            detail["ai_level"] = detail.get("ai_level") or detail.get("level")
            detail["level"] = level
            detail["teacher_adjusted"] = True
            changed = True
        if reason_prefix:
            reason = form.get(f"{reason_prefix}{criterion}", "").strip()
            if reason:
                detail["teacher_reasoning"] = reason
    if changed:
        result["grading"]["subtotals"] = rubric_io.compute_subtotals(
            scores, state.get("rubric_json") or _coerce_rubric(state.get("rubric") or ""))
    return changed


@app.route("/update_scores", methods=["POST"])
def update_scores():
    if not state["grader"]:
        return jsonify({"error": "尚无评分结果"}), 400
    if (state.get("grading_status") or {}).get("phase") == "running":
        return jsonify({"error": "评分进行中，暂不能修改"}), 400
    report_id = request.form.get("report_id")
    for result in state["grader"].results:
        if result["report_id"] == report_id:
            _apply_teacher_levels(result, request.form, "level::", "reason::")
    _export_all()
    return redirect(request.form.get("next") or "/app/results")


@app.route("/approve_report/<path:filename>", methods=["POST"])
def approve_report(filename):
    grader = state["grader"]
    if not grader:
        return jsonify({"error": "尚无评分结果"}), 400
    if filename in _approved_ids():
        return redirect(request.form.get("next") or "/app/results")
    if (state.get("grading_status") or {}).get("phase") == "running":
        return jsonify({"error": "评分进行中，请等待完成后再终审"}), 400
    result = next((r for r in grader.results
                   if r["report_id"] == filename and r.get("status") == "graded"), None)
    if not result:
        return jsonify({"error": "报告未评分或状态异常"}), 400
    if _result_stale(result):
        return jsonify({"error": "评分细则已修改，请重新评分后再终审"}), 400
    text = result.get("_report_text") or ""
    if not text:
        entry = _find_entry(filename)
        text = _report_text(entry) if entry else ""
    grader.finalize_report(filename, text)
    state["approved_ids"] = _approved_ids() | {filename}
    save_approved(state["approved_ids"])
    _export_all()
    return redirect(request.form.get("next") or "/app/results")


# ---------------------------------------------------------------- 配置
@app.route("/settings", methods=["GET", "POST"])
def settings():
    cfg = state["config"]
    ocr_cfg = cfg.setdefault("ocr", {})
    ocr_cfg.setdefault("backend", "paddleocr")
    ocr_cfg.setdefault("paddleocr", {}).setdefault("language", "ch")
    ocr_cfg.setdefault("api", {}).setdefault("api_url", "")
    ocr_cfg["api"].setdefault("api_key", "")
    ocr_cfg["api"].setdefault("api_model", "glm-4v-flash")
    llm_cfg = cfg.setdefault("llm", {})
    llm_cfg.setdefault("backend", "openai")
    for section, defaults in (("openai", {"base_url": "", "api_key": "DEEPSEEK_API_KEY",
                                          "model": "", "max_tokens": 2048}),
                              ("ollama", {"base_url": "http://localhost:11434", "model": ""}),
                              ("custom", {"base_url": "", "api_key": "", "model": ""})):
        target = llm_cfg.setdefault(section, {})
        for key, value in defaults.items():
            target.setdefault(key, value)
    vs = cfg.setdefault("vector_store", {})
    ch = vs.setdefault("chroma", {})
    emb = vs.setdefault("embedding", {})
    emb.setdefault("openai", {})
    for key, default in (("backend", "chroma"),):
        vs.setdefault(key, default)
    for key, default in (("mode", "local"), ("host", ""), ("port", 8000), ("ssl", False), ("token", "")):
        ch.setdefault(key, default)
    emb.setdefault("backend", "default")
    emb["openai"].setdefault("base_url", "")
    emb["openai"].setdefault("api_key", "")
    emb["openai"].setdefault("model", "")
    if request.method == "GET":
        return redirect("/app/settings")
    form = request.form
    backend = form.get("llm_backend", llm_cfg.get("backend", "openai"))
    if backend in BACKENDS:
        llm_cfg["backend"] = backend
    ollama = llm_cfg.setdefault("ollama", {})
    ollama["base_url"] = form.get("ollama_base_url", ollama.get("base_url", ""))
    ollama["model"] = form.get("ollama_model", ollama.get("model", ""))
    custom = llm_cfg.setdefault("custom", {})
    custom["base_url"] = form.get("custom_base_url", custom.get("base_url", ""))
    custom["model"] = form.get("custom_model", custom.get("model", ""))
    if form.get("custom_api_key", "").strip():
        custom["api_key"] = form["custom_api_key"].strip()
    openai = cfg.setdefault("llm", {}).setdefault("openai", {})
    openai["base_url"] = form.get("base_url", openai.get("base_url", ""))
    openai["model"] = form.get("model", openai.get("model", ""))
    if form.get("api_key", "").strip():
        openai["api_key"] = form["api_key"].strip()
    openai["max_tokens"] = _int_or(form.get("max_tokens"), openai.get("max_tokens") or 2048)
    ocr = cfg.setdefault("ocr", {})
    ocr["backend"] = form.get("ocr_backend", ocr.get("backend", "paddleocr"))
    api = ocr.setdefault("api", {})
    api["api_url"] = form.get("ocr_api_url", api.get("api_url", ""))
    if form.get("ocr_api_key", "").strip():
        api["api_key"] = form["ocr_api_key"].strip()
    api["api_model"] = form.get("ocr_api_model", api.get("api_model", ""))
    vs["backend"] = form.get("vs_backend", vs.get("backend", "chroma"))
    ch["mode"] = form.get("vs_mode", ch.get("mode", "local"))
    ch["host"] = form.get("vs_host", ch.get("host", ""))
    ch["port"] = _int_or(form.get("vs_port"), ch.get("port") or 8000)
    ch["ssl"] = form.get("vs_ssl") == "true"
    if form.get("vs_token", "").strip():
        ch["token"] = form["vs_token"].strip()
    emb["backend"] = form.get("emb_backend", emb.get("backend", "default"))
    emb["openai"]["base_url"] = form.get("emb_url", emb["openai"].get("base_url", ""))
    if form.get("emb_key", "").strip():
        emb["openai"]["api_key"] = form["emb_key"].strip()
    emb["openai"]["model"] = form.get("emb_model", emb["openai"].get("model", ""))
    raw_workers = (form.get("grading_workers") or "").strip()
    if raw_workers:
        try:
            grading_workers = int(raw_workers)
        except ValueError:
            return jsonify({"error": "并行数必须是 1-8 的整数"}), 400
        cfg.setdefault("grading", {})["workers"] = max(1, min(8, grading_workers))
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        yaml.safe_dump(cfg, f, allow_unicode=True, sort_keys=False)
    state["config"] = cfg
    return jsonify({"ok": True})


BACKENDS = {"openai": "OpenAI 兼容接口", "ollama": "Ollama（本地）", "custom": "自定义 HTTP"}


def _int_or(value, default):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


# ---------------------------------------------------------------- 模型方案（多套配置一键切换）
def _profile_payload(source):
    """从 llm 配置段（或已存方案）提取方案本体：backend + 三协议块。"""
    import copy

    return {key: copy.deepcopy(source[key]) for key in
            ("backend", "openai", "ollama", "custom") if isinstance(source, dict) and key in source}


def _persist_config():
    with open(CONFIG_PATH, "w", encoding="utf-8") as f:
        yaml.safe_dump(state["config"], f, allow_unicode=True, sort_keys=False)


def _request_name():
    payload = request.get_json(silent=True) or request.form or {}
    return str(payload.get("name", "")).strip()


@app.route("/llm/profile/save", methods=["POST"])
def llm_profile_save():
    """把当前生效的模型配置快照存为命名方案；同名覆盖。"""
    name = _request_name()
    if not name or len(name) > 40:
        return jsonify({"error": "方案名需为 1-40 个字符"}), 400
    llm = state["config"].setdefault("llm", {})
    payload = _profile_payload(llm)
    if not payload:
        return jsonify({"error": "当前配置为空，无法保存"}), 400
    llm.setdefault("profiles", {})[name] = payload
    _persist_config()
    return jsonify({"ok": True, "name": name})


@app.route("/llm/profile/apply", methods=["POST"])
def llm_profile_apply():
    """把方案写入当前生效配置，立即生效（正在跑的评分线程持旧客户端，下一轮生效）。"""
    name = _request_name()
    llm = state["config"].setdefault("llm", {})
    profile = (llm.get("profiles") or {}).get(name)
    if not profile:
        return jsonify({"error": f"方案「{name}」不存在"}), 404
    llm.update(_profile_payload(profile))
    llm["active_profile"] = name
    _persist_config()
    return jsonify({"ok": True, "active": name})


@app.route("/llm/profile/delete", methods=["POST"])
def llm_profile_delete():
    name = _request_name()
    llm = state["config"].setdefault("llm", {})
    profiles = llm.get("profiles") or {}
    if name not in profiles:
        return jsonify({"error": f"方案「{name}」不存在"}), 404
    del profiles[name]
    if llm.get("active_profile") == name:
        llm["active_profile"] = ""  # 只清标记，当前配置不变
    _persist_config()
    return jsonify({"ok": True})


@app.route("/settings/ocr_check", methods=["POST"])
def ocr_check():
    """OCR 自检：用合成图片跑一遍识别，返回识别文本与字符准确率。"""
    result = _run_ocr_check()
    return jsonify(result)


def _run_ocr_check():
    from .utils import ocr_check
    return ocr_check.run((state["config"] or {}).get("ocr", {}))


@app.route("/download/<path:filename>")
def download(filename):
    root = os.path.realpath(session_dir())
    target = os.path.realpath(os.path.join(root, filename))
    if not target.startswith(root + os.sep):
        return "非法路径", 403
    if not os.path.isfile(target):
        return "文件不存在", 404
    return send_file(target, as_attachment=True)


# ---------------------------------------------------------------- SPA API
# 说明：写操作沿用既有表单路由（表单编码 POST 后重新拉取 JSON），
# 这里只补页面数据接口与需要结构化结果的接口，避免两套业务逻辑。

def _api_criteria_uploaded(names):
    return any(os.path.splitext(n)[1].lower() in (".docx", ".pdf", ".txt")
               and not n.startswith("calibration_") for n in names)


@app.route("/")
def spa_home():
    index = os.path.join(UI_DIST, "index.html")
    if not os.path.exists(index):
        return ("前端尚未构建：请执行 cd frontend && npm install && npm run build", 503,
                {"Content-Type": "text/plain; charset=utf-8"})
    return send_file(index)


@app.route("/app")
@app.route("/app/<path:sub>")
def spa(sub=None):
    index = os.path.join(UI_DIST, "index.html")
    if not os.path.exists(index):
        return ("前端尚未构建：请执行 cd frontend && npm install && npm run build", 503,
                {"Content-Type": "text/plain; charset=utf-8"})
    return send_file(index)


@app.route("/favicon.ico")
def favicon():
    return send_file(os.path.join(app.static_folder, "favicon.svg"), mimetype="image/svg+xml")


@app.route("/api/bootstrap")
def api_bootstrap():
    return jsonify(_bootstrap_payload())


@app.route("/api/balance")
def api_balance():
    from .web import balance as balance_store

    force = request.args.get("refresh") == "1"
    events = (state.get("usage") or {}).get("events") or []
    return jsonify(balance_store.refresh(state["config"], events, force=force))


@app.route("/api/balance/reset", methods=["POST"])
def api_balance_reset():
    from .web import balance as balance_store

    balance_store.reset()
    events = (state.get("usage") or {}).get("events") or []
    return jsonify(balance_store.refresh(state["config"], events, force=True))


@app.route("/api/vector")
def api_vector():
    try:
        store = Chroma(config=(state.get("config") or {}).get("vector_store") or {},
                       session_dir=session_dir(), create=False)
        rows = []
        for collection in store.client.list_collections():
            meta = dict(collection.metadata or {})
            rows.append({
                "name": collection.name,
                "count": collection.count(),
                "dimension": _collection_dimension(collection),
                "metadata": _short_meta(meta),
                "kind": ("全文库" if collection.name == "grading_cases"
                         else "维度库" if collection.name.startswith("dim_") else "其他"),
                "label": meta.get("dimension") or collection.name,
            })
        rows.sort(key=lambda row: (row["kind"] != "全文库", row["name"]))
        return jsonify({"collections": rows, "total": len(rows)})
    except Exception as exc:
        return jsonify({"error": f"读取向量库失败：{exc}"}), 500


@app.route("/api/vector/<path:name>")
def api_vector_collection(name):
    try:
        offset = max(0, int(request.args.get("offset", 0)))
        limit = max(1, min(50, int(request.args.get("limit", 20))))
    except ValueError:
        return jsonify({"error": "分页参数无效"}), 400
    try:
        store = Chroma(config=(state.get("config") or {}).get("vector_store") or {},
                       session_dir=session_dir(), create=False)
        collection = store.client.get_collection(name)
    except Exception:
        return jsonify({"error": "collection 不存在"}), 404
    try:
        total = collection.count()
        got = collection.get(limit=limit, offset=offset,
                             include=["documents", "metadatas", "embeddings"])
        ids = got.get("ids") or []
        docs = got.get("documents") or []
        metas = got.get("metadatas") or []
        vectors = got.get("embeddings")
        vectors = list(vectors) if vectors is not None else []
        items = []
        for index, item_id in enumerate(ids):
            doc = docs[index] if index < len(docs) and docs[index] else ""
            vector = vectors[index] if index < len(vectors) else None
            items.append({
                "id": item_id,
                "document": doc[:4000],
                "docChars": len(doc),
                "truncated": len(doc) > 4000,
                "metadata": _short_meta(metas[index] if index < len(metas) else {}),
                "dimension": len(vector) if vector is not None else None,
                "embeddingPreview": [round(float(value), 5) for value in list(vector)[:8]]
                                    if vector is not None else [],
            })
        return jsonify({
            "name": name,
            "label": (collection.metadata or {}).get("dimension") or name,
            "kind": ("全文库" if name == "grading_cases"
                     else "维度库" if name.startswith("dim_") else "其他"),
            "total": total,
            "offset": offset,
            "limit": limit,
            "dimension": _collection_dimension(collection),
            "metadata": _short_meta(dict(collection.metadata or {})),
            "items": items,
        })
    except Exception as exc:
        return jsonify({"error": f"读取 collection 失败：{exc}"}), 500


@app.route("/api/vector/clear", methods=["POST"])
def api_vector_clear():
    if _busy():
        return jsonify({"error": "有任务正在进行，请等评分/解析结束后再清空"}), 409
    try:
        store = Chroma(config=(state.get("config") or {}).get("vector_store") or {},
                       session_dir=session_dir(), create=False)
        names = [collection.name for collection in store.client.list_collections()]
        for name in names:
            store.client.delete_collection(name)
        if state.get("grader"):
            state["grader"].invalidate_vector_store()
        return jsonify({"ok": True, "deleted": len(names)})
    except Exception as exc:
        return jsonify({"error": f"清空向量库失败：{exc}"}), 500


def _collection_dimension(collection):
    try:
        got = collection.get(limit=1, include=["embeddings"])
        vectors = got.get("embeddings")
        if vectors is not None and len(vectors):
            return len(vectors[0])
    except Exception:
        pass
    return None


def _short_meta(value, limit=800):
    if isinstance(value, str):
        return value if len(value) <= limit else value[:limit] + "…"
    if isinstance(value, dict):
        return {key: _short_meta(item, limit) for key, item in value.items()}
    if isinstance(value, list):
        return [_short_meta(item, limit) for item in value[:20]]
    return value


@app.route("/api/usage")
def api_usage():
    from .web import usage as usage_store

    try:
        days = max(0, min(3650, int(request.args.get("days", 1))))
    except ValueError:
        days = 1
    model = (request.args.get("model") or "").strip() or None
    return jsonify(usage_store.payload(session_dir(), days, model))


@app.route("/api/rubric")
def api_rubric():
    rubric = state.get("rubric_json")
    uploads = os.path.join(session_dir(), "uploads")
    names = os.listdir(uploads) if os.path.isdir(uploads) else []
    external = state.get("external") or {}
    counts = {name: len(values or {})
              for name, values in (external.get("scores") or {}).items()}
    criteria = [{**c, "can_exclude": rubric_io.is_non_report(c),
                 "external_count": counts.get(c.get("name"), 0)}
                for c in ((rubric or {}).get("criteria") or [])]
    return jsonify({
        "rubric": {**(rubric or {}), "criteria": criteria} if rubric else rubric,
        "criteria": criteria,
        "warnings": state.get("rubric_warnings") or [],
        "hasMarkdown": os.path.exists(os.path.join(session_dir(), "rubrics", "rubric_latest.md")),
        "uploaded": _api_criteria_uploaded(names),
        "docChars": len(state.get("criteria_doc") or ""),
        "experimentChars": len(state.get("experiment_doc") or ""),
        "mode": "pair" if (state.get("experiment_doc") or "").strip() else "single",
        "external": external_scores.summary(external),
    })


@app.route("/api/reports")
def api_reports():
    graded = _graded_ids()
    approved = _approved_ids()
    rows = []
    for r in state["extracted_reports"]:
        if r.get("status") != "ok":
            continue
        meta = r["meta"]
        rows.append({
            "filename": meta["filename"],
            "student": meta.get("student") or {},
            "pages": meta.get("pages"),
            "paragraphs": meta.get("paragraphs"),
            "images": meta.get("image_count") or 0,
            "scanned": bool(meta.get("is_scanned")),
            "graded": meta["filename"] in graded,
            "approved": meta["filename"] in approved,
            "hasOriginal": _origin_path(meta["filename"]) is not None,
        })
    return jsonify({"reports": rows, "fields": _student_fields(), "downloads": _download_paths()})


@app.route("/api/report/<path:filename>")
def api_report(filename):
    _apply_external()
    entry = _find_entry(filename)
    if not entry or entry.get("status") != "ok":
        return jsonify({"error": "报告不存在"}), 404
    text = _report_text(entry)
    result = next((r for r in (state["grader"].results if state["grader"] else [])
                   if r["report_id"] == filename), None)
    blocks = _build_dim_blocks(state.get("rubric_json") or _coerce_rubric(state.get("rubric") or ""),
                               text, result, entry.get("images") or ())
    return jsonify({
        "filename": filename,
        "student": entry["meta"].get("student") or {},
        "fields": _student_fields(),
        "blocks": blocks,
        "result": result,
        "hasOriginal": _origin_path(filename) is not None,
    })


@app.route("/api/select")
def api_select():
    ctx = _select_context()
    return jsonify({
        "groups": [{"label": g["label"], "multi": g["multi"],
                    "reports": [{"filename": r["meta"]["filename"],
                                 "student": r["meta"].get("student") or {},
                                 "pages": r["meta"].get("pages"),
                                 "paragraphs": r["meta"].get("paragraphs"),
                                 "images": r["meta"].get("image_count") or 0,
                                 "scanned": bool(r["meta"].get("is_scanned")),
                                 "graded": r["meta"]["filename"] in ctx["graded_ids"],
                                 "stale": r["meta"]["filename"] in ctx["stale_ids"],
                                 "status": (ctx["statuses"].get(r["meta"]["filename"]) or {}).get("status", "")}
                                for r in g["reports"]]} for g in ctx["groups"]],
        "hint": ctx["hint"],
        "hasRubric": ctx["has_rubric"],
        "rubricConfirmed": ctx.get("rubric_confirmed"),
        "error": ctx["error"],
        "total": len(ctx["reports"]),
    })


@app.route("/api/results")
def api_results():
    _apply_external()
    results = state["grader"].results if state["grader"] else []
    approved = _approved_ids()
    running = (state.get("grading_status") or {}).get("phase") == "running"
    students, suspicious, _missing = _classify_students()
    counts = {
        "all": len(results),
        "graded": len([r for r in results if r.get("status") == "graded"]),
        "injection": len([r for r in results if r.get("status") == "flagged_injection"]),
        "blocked": len([r for r in results if r.get("status") == "blocked_safety"]),
        "failed": len([r for r in results if r.get("status") in ("llm_failed", "detection_failed",
                                                                 "audit_truncated")]),
        "plagiarism": len([r for r in results if r.get("similarity_warnings")]),
        "unapproved": len([r for r in results if r["report_id"] not in approved]),
    }
    return jsonify({
        "results": results,
        "counts": counts,
        "approvedIds": sorted(approved),
        "students": students,
        "staleIds": [r["report_id"] for r in results if _result_stale(r)],
        "running": running,
        "downloads": _download_paths(),
        "agreement": _agreement(),
        "fields": _student_fields(),
        "hasReports": bool(state["extracted_reports"]),
        "hasRubric": bool(state.get("rubric_json") or state.get("rubric")),
    })


@app.route("/api/roster")
def api_roster():
    roster = load_roster()
    _students, suspicious, missing = _classify_students()
    return jsonify({
        "roster": roster,
        "reportCount": len([r for r in state["extracted_reports"] if r.get("status") == "ok"]),
        "suspicious": suspicious,
        "missing": missing,
        "mostlyUnmatched": bool(roster and suspicious and len(suspicious) * 2 >= max(1, len(
            [r for r in state["extracted_reports"] if r.get("status") == "ok"]))),
    })


@app.route("/api/settings")
def api_settings():
    cfg = state["config"]
    llm_cfg = cfg.get("llm") or {}
    profiles = []
    for name, profile in (llm_cfg.get("profiles") or {}).items():
        backend = profile.get("backend") or "openai"
        profiles.append({"name": name, "backend": backend,
                         "model": ((profile.get(backend) or {}).get("model")) or ""})
    return jsonify({
        "llm": llm_cfg,
        "ocr": cfg.get("ocr") or {},
        "vectorStore": cfg.get("vector_store") or {},
        "grading": cfg.get("grading") or {},
        "profiles": profiles,
        "activeProfile": llm_cfg.get("active_profile") or "",
        "backends": BACKENDS,
        "backendLabel": BACKENDS.get(llm_cfg.get("backend", "openai"), ""),
    })


@app.route("/api/mcp")
def api_mcp():
    return jsonify({
        "servers": mcp_store.load_servers(),
        "mount": state.get("mcp") or load_mcp_mount(),
        "status": state.get("mcp_status") or [],
        "referenced": state.get("mcp_refs") or "",
    })


@app.route("/api/calibration")
def api_calibration():
    calib = state.get("calibration")
    if not calib or not state.get("rubric_json"):
        return jsonify({"ready": False})
    rows = []
    for c in state["rubric_json"]["criteria"]:
        paras = [calib["chunks"][int(aid[1:])] for aid, a in calib["assignments"].items()
                 if (a or {}).get("criterion") == c["name"] and aid[1:].isdigit()
                 and int(aid[1:]) < len(calib["chunks"])]
        rows.append({"criterion": c, "paragraphs": paras})
    mounted = set(calib["assignments"])
    return jsonify({
        "ready": True,
        "filename": calib["filename"],
        "rows": rows,
        "rest": [p for i, p in enumerate(calib["chunks"]) if f"p{i}" not in mounted],
    })


@app.route("/api/calibration/submit", methods=["POST"])
def api_calibration_submit():
    calib = state.get("calibration")
    if not calib:
        return jsonify({"error": "没有待校准的报告"}), 400
    payload = request.get_json(silent=True) or {}
    levels = payload.get("levels") or {}
    reasons = payload.get("reasons") or {}
    teacher = {name: {"level": level, "evidence": "", "reasoning": str(reasons.get(name, "")).strip()}
               for name, level in levels.items() if level in VALID_LEVELS}
    grader = Grader(state["rubric_json"], llm_client=llm(state["config"]), config=state["config"],
                    similarity_scope=state.get("session"),
                    kb_session_dir=os.path.join(session_dir(), "calibration_kb"))
    outcome = grader.grade_single("calibration", calib["text"], references=_mcp_references("校准"))
    result = outcome["result"]
    if result.get("status") != "graded":
        detail = (result.get("injection_details") or {}).get("evidence") or result.get("error") or ""
        return jsonify({"error": f"AI 评分未完成（{result.get('status')}）：{detail}", "status": result.get("status")})
    analysis = RubricGenerator(llm_client=llm(state["config"])).calibrate(
        state.get("rubric", ""), calib["text"], teacher, result["grading"])
    return jsonify({"aiGrading": result["grading"], "analysis": analysis})


if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1", host="127.0.0.1", port=5000)
