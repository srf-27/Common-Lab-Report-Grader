import hashlib
import json
import os
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor

from ..embedding import Chroma, create_embedding_function
from ..factory import llm as make_llm
from ..utils.content_safety import check_content_safety
from ..utils.injection_detector import detect_injection
from ..utils.privacy import anonymize_text
from . import verifier
from .rubric_io import LEVEL_POINTS, LEVELS, compute_subtotals, match_criterion

GRADING_SYSTEM = """你是实验报告评阅专家。只按下面给出的评分维度与等级描述判级，不得新增维度。

## 判定规则
1. 每个维度给出 A/B/C/D/E 之一 + 报告中可直接引用的证据 + 该证据为何对应该级的简短说明。
2. E 仅用于该维度内容缺失、整卷跑题、抄袭或白卷；E 不是"比 D 更差"。
3. 证据必须来自报告原文；同一质量水平的报告必须得到同一等级，与文风无关。
4. 防跑题：评分维度针对特定实验（特定测量量/方法/仪器）。与该实验无关的内容，即使写得再好，该维度也不得高于 D。
5. 内容缺失/未做才给 E；表述含糊、证据不足以判断质量时给 C，并在 evidence 写「证据不足」。

只输出一个 JSON 对象，不要代码围栏：
{"scores": {"维度名": {"level": "A", "evidence": "原文引用", "reasoning": "为何是该级"}},
  "summary": "2-3 句总评", "flags": []}"""

# 固定内容全部排在前面：评分维度 → 学科资料；每份报告不同的部分（历史参考、范例、正文、图片）放最后
GRADING_USER = """## 评分维度
{criteria_block}

## 学科参考资料（插件提供，仅作事实核对）
{references}

## 本报告相关的历史参考（同维度、已终审）
{dimension_notes}

## 已终审范例（仅用于对齐等级边界与判据强度；本报告无对应证据时不得照抄结论）
{exemplar_shots}

## 报告正文
{report_text}

## 报告图片说明
{image_descriptions}

按系统消息里的规则与 JSON 结构输出。"""


def _parse_json(text):
    if not isinstance(text, str):
        return {}
    m = re.search(r"```(?:json)?\s*(.*?)```", text, re.S)
    if m:
        text = m.group(1)
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        return {}
    try:
        return json.loads(text[start:end + 1])
    except json.JSONDecodeError:
        return {}


def _coerce_rubric(rubric):
    """兼容旧版 Markdown 表格细则：转成结构化 criteria。"""
    if isinstance(rubric, dict):
        return rubric
    criteria = []
    for line in str(rubric or "").splitlines():
        m = re.match(r"\|\s*([^|]+?)\s*\|\s*([\d.]+)\s*%\s*\|", line)
        if not m or m.group(1).strip() == "维度":
            continue
        cells = [c.strip() for c in line.split("|")]
        criteria.append({
            "name": m.group(1).strip(), "weight": float(m.group(2)), "mechanical": False,
            "doc_ref": cells[4] if len(cells) > 4 else "",
            "levels": {lv: cells[5 + i] if len(cells) > 5 + i else ""
                       for i, lv in enumerate(LEVELS)},
            "user_edited": [], "rules": None,
        })
    return {"criteria": criteria}


class Grader:
    def __init__(self, rubric, llm_client=None, config=None, output_dir=None,
                 similarity_scope=None, kb_session_dir=None):
        self.rubric = _coerce_rubric(rubric)
        self.llm = llm_client or make_llm(config)
        self.config = config or {}
        self.output_dir = output_dir or os.path.join(os.path.dirname(__file__), "..", "output", "grading")
        os.makedirs(self.output_dir, exist_ok=True)
        self.results = []
        self._results_lock = threading.RLock()
        self._kb_lock = threading.RLock()
        self._cache_lock = threading.RLock()
        self._embedding_lock = threading.RLock()
        self._embedding_cache = {}
        try:
            self.parallel_checks = int(os.environ.get("ERG_GRADING_WORKERS", "3")) > 1
        except ValueError:
            self.parallel_checks = True
        self.session_scope = similarity_scope or "default"
        vector_cfg = self.config.get("vector_store") or {}
        self._vector_cfg = vector_cfg
        self._kb_session_dir = kb_session_dir
        self.kb = Chroma(config=vector_cfg, session_dir=kb_session_dir)  # 全文相似度库
        self.embedding_function = create_embedding_function(vector_cfg.get("embedding"))
        grading_cfg = self.config.get("grading") or {}
        self.similarity_threshold = grading_cfg.get("similarity_threshold", 0.92)
        self.similarity_reject = grading_cfg.get("similarity_reject", 0.97)
        self.dimension_collections = {}
        self.remembered = None
        self.verifier_code = None
        self.verifier_error = ""
        self._checks_cache = None

    # ---------- 检测结果缓存 ----------
    def _checks_path(self):
        return os.path.join(self.output_dir, "checks_cache.json")

    def _load_checks(self):
        with self._cache_lock:
            if self._checks_cache is None:
                try:
                    with open(self._checks_path(), encoding="utf-8") as f:
                        self._checks_cache = json.load(f)
                except Exception:
                    self._checks_cache = {}
            return self._checks_cache

    def _save_checks(self):
        cache = self._checks_cache or {}
        if len(cache) > 1000:
            keep = sorted(cache.items(), key=lambda kv: kv[1].get("at", 0))[-1000:]
            cache = dict(keep)
            self._checks_cache = cache
        path = self._checks_path()
        try:
            os.makedirs(os.path.dirname(path), exist_ok=True)
            tmp = path + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(cache, f, ensure_ascii=False)
            os.replace(tmp, path)
        except Exception:
            pass

    def _cached_check(self, kind, text):
        """同一份正文重评时跳过重复检测：正文不变，检测结论就不会变。"""
        digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
        key = f"{kind}:{digest}"
        cache = self._load_checks()
        if key in cache:
            return {k: v for k, v in cache[key].items() if k != "at"}
        if kind == "injection":
            value = detect_injection(text, self.llm)
        else:
            value = check_content_safety(text, self.llm)
        # 检测失败或只审了前半段都不缓存，下次重试
        if not value.get("detection_failed") and not value.get("truncated"):
            with self._cache_lock:
                cache[key] = {**value, "at": time.time()}
                self._save_checks()
        return value

    # ---------- RAG ----------
    def _ensure_kb(self):
        if self.kb is None:
            self.kb = Chroma(config=self._vector_cfg, session_dir=self._kb_session_dir)
            self.dimension_collections = {}
            self.remembered = None

    def invalidate_vector_store(self):
        """清空向量库后丢弃旧 collection 句柄，下次终审再按需重建。"""
        self.kb = None
        self.dimension_collections = {}
        self.remembered = None

    def _query_vector(self, text):
        """同一份正文的 query embedding 只算一次，维度参考不再逐维重新嵌入。"""
        key = hashlib.sha256(text.encode("utf-8")).hexdigest()
        with self._embedding_lock:
            if key in self._embedding_cache:
                return self._embedding_cache[key]
        try:
            vector = self.kb.collection._embed([text], is_query=True)[0]
        except Exception:
            vector = None
        with self._embedding_lock:
            if len(self._embedding_cache) > 200:
                self._embedding_cache.clear()
            self._embedding_cache[key] = vector
        return vector

    def _query_collection(self, collection, text, n_results):
        vector = self._query_vector(text)
        if vector is not None:
            try:
                return collection.query(query_embeddings=[vector], n_results=n_results)
            except Exception:
                pass
        return collection.query(query_texts=[text], n_results=n_results)

    def _dimension_collection(self, name, create=True):
        """每个维度一个库；create=False 只读，rubric 确认前不会凭空建库。"""
        with self._kb_lock:
            if name in self.dimension_collections:
                return self.dimension_collections[name]
            key = "dim_" + hashlib.sha1(name.encode("utf-8")).hexdigest()[:12]
            if not create:
                try:
                    return self.kb.client.get_collection(
                        name=key, embedding_function=self.embedding_function)
                except Exception:
                    return None
            self.dimension_collections[name] = self.kb.client.get_or_create_collection(
                name=key, metadata={"hnsw:space": "cosine", "dimension": name},
                embedding_function=self.embedding_function)
            return self.dimension_collections[name]

    def _dimension_references(self, name, text, n_results=2):
        collection = self._dimension_collection(name, create=False)
        if collection is None:
            return ""
        with self._kb_lock:
            total = collection.count()
            if total == 0 or not text.strip():
                return ""
            hits = self._query_collection(collection, text, min(n_results, total))
        return "\n".join(
            f"- （相似度 {1 - distance:.2f}，历史等级 {meta.get('level') or '未定'}）{doc[:200]}"
            for doc, meta, distance in zip(hits["documents"][0], hits["metadatas"][0],
                                           hits["distances"][0]))

    def _known_ids(self):
        if self.remembered is None:
            self.remembered = set(self.kb.collection.get(include=[]).get("ids") or [])
        return self.remembered

    def _check_similarity(self, report_id, text):
        with self._kb_lock:
            total = self.kb.count()
            if total == 0 or not text.strip():
                return []
            hits = self._query_collection(self.kb.collection, text, min(total, 10))
        warnings = []
        for matched, distance in zip(hits["ids"][0], hits["distances"][0]):
            if matched == report_id:
                continue
            similarity = 1 - distance
            if similarity >= self.similarity_threshold:
                warnings.append({"matched_id": matched, "similarity": round(similarity, 4),
                                 "level": "rejected" if similarity >= self.similarity_reject else "warning"})
        return warnings

    def _remember_report(self, report_id, text, grading=None):
        if not text.strip():
            return False
        with self._kb_lock:
            known = self._known_ids()
            already = report_id in known
        if already:
            if grading:
                metadata = {"report_id": report_id, "scope": self.session_scope,
                            "grading_json": anonymize_text(
                                json.dumps(grading, ensure_ascii=False), self.llm)[0]}
                with self._kb_lock:
                    self.kb.collection.update(ids=[report_id], metadatas=[metadata])
            return False
        anon_text = anonymize_text(text, self.llm)[0]
        metadata = {"report_id": report_id, "scope": self.session_scope}
        if grading:
            metadata["grading_json"] = anonymize_text(json.dumps(grading, ensure_ascii=False), self.llm)[0]
        with self._kb_lock:
            if report_id in self._known_ids():
                if grading:
                    self.kb.collection.update(ids=[report_id], metadatas=[metadata])
                return False
            self.kb.add(ids=[report_id], documents=[anon_text], metadatas=[metadata])
            self.remembered.add(report_id)
            return True

    def _exemplar_shots(self, query_text, n_cases=2, max_chars=1200):
        """取 1-2 份得分最高的已终审报告（正文节选 + 其各维度判定）作为提示词 shots。"""
        with self._kb_lock:
            total = self.kb.count()
            if total == 0 or not query_text.strip():
                return ""
            hits = self._query_collection(self.kb.collection, query_text, min(total, 5))
        cases = []
        for document, metadata in zip(hits["documents"][0], hits["metadatas"][0]):
            grading = _parse_json((metadata or {}).get("grading_json") or "")
            scores = grading.get("scores") or {}
            if not scores:
                continue
            points = sum(LEVEL_POINTS.get((entry or {}).get("level"), 0) for entry in scores.values())
            cases.append((points, document, scores))
        cases.sort(key=lambda case: case[0], reverse=True)
        shots = []
        for index, (_, document, scores) in enumerate(cases[:n_cases], 1):
            judgements = "\n".join(
                f"- {name}：{entry.get('level', '')}（依据：{(entry.get('evidence') or '')[:80]}）"
                for name, entry in scores.items())
            shots.append(f"### 范例 {index}\n【全文节选】\n{document[:max_chars]}\n"
                         f"【该范例各维度判定】\n{judgements}")
        return "\n\n".join(shots)

    def _anonymize_many(self, items):
        """一次 LLM 调用脱敏多条文本：用固定分隔符拼接，替换后原样切回。"""
        if not items:
            return []
        separator = "\n\n<<<SPLIT>>>\n\n"
        joined = separator.join(items)
        scrubbed = anonymize_text(joined, self.llm)[0]
        parts = scrubbed.split(separator)
        if len(parts) == len(items):
            return parts
        return [anonymize_text(item, self.llm)[0] for item in items]  # 兜底：逐条

    def _remember_dimension_chunks(self, result):
        """每个学生每个维度只存一条记录：把挂到该维度的 chunk 合并成一个大块再嵌入。"""
        chunks = result.get("_report_chunks") or []
        assignments = result.get("text_assignments") or {}
        images = result.get("images") or []
        scores = (result.get("grading") or {}).get("scores") or {}
        groups = {}
        for chunk_id, assignment in assignments.items():
            criterion = (assignment or {}).get("criterion")
            index = int(chunk_id[1:]) if chunk_id.startswith("p") and chunk_id[1:].isdigit() else -1
            if not criterion or not 0 <= index < len(chunks):
                continue
            group = groups.setdefault(criterion, {"chunks": [], "images": []})
            if chunks[index].strip():
                group["chunks"].append((index, chunks[index]))
        for filename, assignment in (result.get("image_assignments") or {}).items():
            criterion = (assignment or {}).get("criterion")
            image = next((im for im in images if im["filename"] == filename), None)
            if not criterion or image is None:
                continue
            groups.setdefault(criterion, {"chunks": [], "images": []})["images"].append(image)
        entries = []
        for criterion, group in groups.items():
            parts = [text for _, text in sorted(group["chunks"], key=lambda item: item[0])]
            parts += [f"[图] {image['filename']} {image.get('context') or ''}".strip()
                      for image in group["images"]]
            document = "\n\n".join(part for part in parts if str(part).strip())
            if not document:
                continue
            entry = scores.get(criterion) or {}
            entries.append({
                "dimension": criterion,
                "document": document,
                "level": entry.get("level", ""),
                "evidence": str(entry.get("evidence") or "")[:500],
                "chunk_count": len(group["chunks"]),
                "image_count": len(group["images"]),
            })
        if not entries:
            return 0
        scrubbed = self._anonymize_many([entry["document"] for entry in entries])
        for entry, document in zip(entries, scrubbed):
            metadata = {"report_id": result["report_id"], "dimension": entry["dimension"],
                        "level": entry["level"], "evidence": entry["evidence"],
                        "chunk_count": entry["chunk_count"], "image_count": entry["image_count"]}
            collection = self._dimension_collection(entry["dimension"])
            item_id = f"{result['report_id']}__dim_" + hashlib.sha1(
                entry["dimension"].encode("utf-8")).hexdigest()[:10]
            with self._kb_lock:
                try:
                    collection.delete(where={"$and": [{"report_id": result["report_id"]},
                                                      {"dimension": entry["dimension"]}]})
                except Exception:
                    pass
                # 同一维度的大块用同一个 id：重复终审时覆盖，不会越积越多
                collection.upsert(ids=[item_id], documents=[document], metadatas=[metadata])
        return len(entries)

    # ---------- 提示词组装 ----------
    def _criteria_block(self, criteria):
        """只输出维度本身：这部分在同一门课的所有报告里完全一致，放在前面对前缀缓存友好。"""
        blocks = []
        for c in criteria:
            lv = c.get("levels") or {}
            blocks.append(
                f"### {c['name']}（权重 {c.get('weight')}%，标准文档位置：{c.get('doc_ref', '未标注')}）\n"
                + "\n".join(f"- {name}: {lv.get(name, '')}" for name in LEVELS))
        return "\n\n".join(blocks)

    def _dimension_notes(self, criteria, query_text):
        """每份报告不同的维度历史片段：放在提示词尾部，避免打断可缓存的前缀。"""
        parts = []
        for c in criteria:
            references = self._dimension_references(c["name"], query_text) if query_text else ""
            if references:
                parts.append(f"### {c['name']}\n{references}")
        return "\n\n".join(parts) or "（暂无）"

    # ---------- 机械判决 ----------
    def _verifier(self):
        """机械判决项的裁判代码：首次使用时由 LLM 生成，并随 rubric 保存供复核。"""
        if self.verifier_code is None:
            mechanical = [c for c in self.rubric.get("criteria", [])
                          if c.get("mechanical") and not c.get("excluded")]
            self.verifier_code = self.rubric.get("verifier_code") or ""
            if mechanical and not self.verifier_code:
                try:
                    self.verifier_code = verifier.generate(mechanical, self.llm)
                    self.rubric["verifier_code"] = self.verifier_code
                except Exception as e:
                    self.verifier_error = f"verifier 生成失败：{e}"
                    self.verifier_code = ""
        return self.verifier_code

    def _verify_mechanical(self, text):
        scores, fallback = {}, []
        mechanical = [c for c in self.rubric.get("criteria", [])
                      if c.get("mechanical") and not c.get("excluded")]
        if not mechanical:
            return scores, fallback
        results = {}
        code = self._verifier()
        if code:
            try:
                results = verifier.run(code, text)
            except Exception as e:
                self.verifier_error = f"verifier 执行失败：{e}"
                results = {}
        for c in mechanical:
            entry = results.get(c["name"])
            if not isinstance(entry, dict) or entry.get("passed") is None:
                rule_result = verifier.verify(text, c.get("rules"))
                if rule_result["verdict"] != "ok":
                    fallback.append(c)  # 无人可判 → 交回 LLM
                    continue
                entry = {"passed": rule_result["passed"], "detail": rule_result["detail"]}
            scores[c["name"]] = {
                "level": "A" if entry.get("passed") else "D",
                "evidence": entry.get("detail", ""),
                "reasoning": "程序化判决（LLM 生成的 verifier）",
                "mechanical": True,
            }
        return scores, fallback

    # ---------- 内容归维度 ----------
    def assign_items(self, entries, shots_query=""):
        """把文本片段与图片说明一次性归入评分维度；LLM 失败返回 {}（内容落入「未挂载」）。"""
        if not entries:
            return {}
        names = [c["name"] for c in self.rubric.get("criteria", [])]
        items = [{"id": e["id"], "text": (e.get("text") or "")[:800]} for e in entries]
        shots = self._exemplar_shots(shots_query or " ".join(i["text"] for i in items),
                                     n_cases=1, max_chars=600)
        listing = "\n".join(f"{index + 1}. {name}" for index, name in enumerate(names))
        prompt = (
            "把下面每一条内容归入最相关的评分维度。评分维度（criterion 必须原样复制下面的名称，"
            "不要改写、缩写或合并；也可以直接填序号）：\n" + listing
            + "\n规则：每条只能归入一个最相关维度。表格、数据记录、公式推导、计算过程、实验步骤都是正文，"
              "必须归入对应维度；只有不含实质内容的封面/页眉元信息（实验名称、姓名、学号、班级、日期）才填 null。"
              "拿不准时归入最相关维度，禁止因为内容是表格或公式而填 null。"
              "\n逐条快速判断即可：不要复述条目内容，不要写分析过程、理由或解释，"
              "直接在 JSON 里给出每条的唯一维度名。每一项 id 都必须出现一次。"
              "\n只输出 JSON：{\"assignments\": [{\"id\": \"...\", \"criterion\": \"...\"}]}"
            + (f"\n参考范例（已终审，仅供对齐维度粒度）：\n{shots}" if shots else "")
            + "\n待归类内容（JSON）：" + json.dumps(items, ensure_ascii=False)
        )
        try:
            data = _parse_json(self.llm.chat([{"role": "user", "content": prompt}], temperature=0.1))
        except Exception:
            return {}
        result = {}
        for assignment in data.get("assignments", []):
            criterion = assignment.get("criterion")
            if isinstance(criterion, (int, float)) or (
                    isinstance(criterion, str) and criterion.strip().isdigit()):
                index = int(criterion) - 1
                matched = names[index] if 0 <= index < len(names) else None
            else:
                matched = match_criterion(criterion, names)
            if assignment.get("id") and matched:
                result[assignment["id"]] = {"criterion": matched,
                                            "reason": assignment.get("reason", "")}
        if not result:
            result = self._assign_by_keywords(items, names)
        return result

    @staticmethod
    def _assign_by_keywords(items, names):
        """模型没给出可用归类时的兜底：按维度名的字符二元组重合度挂载。"""
        grams = {name: {name[i:i + 2] for i in range(len(name) - 1)} for name in names}
        result = {}
        for item in items:
            text = item.get("text") or ""
            best, score = None, 0
            for name in names:
                hit = sum(1 for gram in grams[name] if gram in text)
                if hit > score:
                    best, score = name, hit
            if best and score >= 2:
                result[item["id"]] = {"criterion": best, "reason": "关键词兜底"}
        return result

    # ---------- 单份评分 ----------
    def grade_single(self, report_id, text, image_descriptions=None, references=""):
        # 注入与内容安全互不依赖：并行发，正常报告少一次模型往返
        if self.parallel_checks:
            with ThreadPoolExecutor(max_workers=2, thread_name_prefix="check") as pool:
                injection_future = pool.submit(self._cached_check, "injection", text)
                safety_future = pool.submit(self._cached_check, "safety", text)
                injection = injection_future.result()
                safety = safety_future.result()
        else:
            injection = self._cached_check("injection", text)
            safety = None
        if injection["is_injection"]:
            return self._finish({"report_id": report_id, "status": "flagged_injection",
                                 "injection_details": injection, "graded_at": time.time()}, False)
        if injection.get("detection_failed") or injection.get("truncated"):
            # 检测器挂了或只审了前 2 万字，都不能静默放行，转人工审查
            status = "detection_failed" if injection.get("detection_failed") else "audit_truncated"
            return self._finish({"report_id": report_id, "status": status,
                                 "injection_details": injection, "graded_at": time.time()}, False)

        if safety is None:
            safety = self._cached_check("safety", text)
        if safety["verdict"] == "blocked":
            return self._finish({"report_id": report_id, "status": "blocked_safety",
                                 "safety_details": safety, "graded_at": time.time()}, False)

        similarity_warnings = self._check_similarity(report_id, text)
        self._remember_report(report_id, text)
        scores, fallback = self._verify_mechanical(text)
        semantic = [c for c in self.rubric.get("criteria", [])
                    if c["name"] not in scores and not c.get("excluded")] + fallback
        if not semantic:
            # 所有维度都由机械判决或成绩单提供，不再调用模型
            grading = {"scores": dict(scores), "summary": "全部维度已由程序或成绩单给出",
                       "flags": [], "subtotals": compute_subtotals(scores, self.rubric)}
            return self._finish({"report_id": report_id, "status": "graded", "grading": grading,
                                 "graded_at": time.time(), "similarity_warnings": similarity_warnings,
                                 "safety_warnings": safety.get("warned", [])}, True)
        # 系统消息与维度块在同一门课内完全一致：稳定前缀越长，供应商的前缀缓存命中越多
        user_prompt = GRADING_USER.format(
            criteria_block=self._criteria_block(semantic),
            dimension_notes=self._dimension_notes(semantic, text),
            exemplar_shots=self._exemplar_shots(text) or "（暂无已终审范例）",
            references=references or "（无插件参考）",
            report_text=text[:8000],
            image_descriptions=image_descriptions or "（本报告无图片）",
        )
        try:
            grading = _parse_json(self.llm.chat(
                [{"role": "system", "content": GRADING_SYSTEM},
                 {"role": "user", "content": user_prompt}], temperature=0.1))
        except Exception as e:
            grading = {"error": str(e)[:200]}
        # 完全没有可解析的分数时不伪造 C：标记失败，交人工
        if not isinstance(grading.get("scores"), dict) or not grading["scores"]:
            return self._finish({"report_id": report_id, "status": "llm_failed",
                                 "error": grading.get("error") or "模型未返回可解析的 JSON",
                                 "graded_at": time.time()}, True)

        grading["scores"] = self._normalize_scores(grading["scores"], semantic, scores)
        grading["subtotals"] = compute_subtotals(grading["scores"], self.rubric)
        result = {"report_id": report_id, "status": "graded", "grading": grading,
                  "graded_at": time.time(), "similarity_warnings": similarity_warnings,
                  "safety_warnings": safety.get("warned", [])}
        if fallback or self.verifier_error:
            result["warnings"] = ([f"机械判决回退 LLM：{'、'.join(c['name'] for c in fallback)}"]
                                  if fallback else []) + ([self.verifier_error] if self.verifier_error else [])
        return self._finish(result, True)

    def _normalize_scores(self, raw_scores, semantic, mechanical):
        """LLM 会自创简短维度名：先归一化回 rubric 名单，否则重复维度按各自权重重复计分。"""
        names = [c["name"] for c in self.rubric.get("criteria", [])]
        normalized = {name: raw_scores[name] for name in names if name in raw_scores}
        for name, entry in raw_scores.items():
            match = match_criterion(name, [v for v in names if v not in normalized])
            if match:
                normalized[match] = entry
        for name, entry in mechanical.items():
            normalized.setdefault(name, entry)
        for c in semantic:
            entry = normalized.get(c["name"])
            if not isinstance(entry, dict) or entry.get("level") not in LEVELS:
                normalized[c["name"]] = {"level": "C", "evidence": "",
                                         "reasoning": "LLM 未返回该维度，默认 C，请人工复核"}
        ordered = {name: normalized[name] for name in names if name in normalized}
        for name, entry in normalized.items():
            ordered.setdefault(name, entry)
        return ordered

    def _finish(self, result, needs_review):
        with self._results_lock:
            self.results = [r for r in self.results if r["report_id"] != result["report_id"]]
            self.results.append(result)
        return {"result": result, "needs_review": needs_review}

    def record_failure(self, report_id, error):
        with self._results_lock:
            self.results = [r for r in self.results if r["report_id"] != report_id]
            self.results.append({"report_id": report_id, "status": "llm_failed",
                                 "error": str(error)[:300], "graded_at": time.time()})

    def finalize_report(self, report_id, text=None):
        """单份报告终审入 RAG：全文进全文库，维度内容汇总进对应维度库。"""
        self._ensure_kb()
        result = next((r for r in self.results if r["report_id"] == report_id), None)
        if result is None:
            return None
        if text and not result.get("_report_text"):
            result["_report_text"] = text
        self._remember_report(report_id, result.get("_report_text", ""), result.get("grading"))
        self._remember_dimension_chunks(result)
        return result

    # ---------- 导出 ----------
    def export_results(self):
        path = os.path.join(self.output_dir, "grading_results.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(self.results, f, ensure_ascii=False, indent=2)
        return path

    def export_excel(self, roster=None, students=None, suspicious=None, missing=None):
        import openpyxl
        path = os.path.join(self.output_dir, "grading_summary.xlsx")
        students = students or {}
        stu_cols = (roster or {}).get("headers") or ["班级", "学院", "学号", "姓名"]
        dims = [(c.get("name", ""), float(c.get("weight") or 0))
                for c in (self.rubric or {}).get("criteria", []) if c.get("name")]
        if not dims:
            seen = []
            for r in self.results:
                for k in ((r.get("grading") or {}).get("scores") or {}):
                    if k not in seen:
                        seen.append(k)
            dims = [(k, 0.0) for k in seen]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "评分汇总"
        dim_cols = [f"{name}({w:g}%)" for name, w in dims]
        ws.append(stu_cols + ["报告ID", "状态", "程序化小计", "语义小计", "总分"] + dim_cols)
        for r in self.results:
            row = [students.get(r["report_id"], {}).get(c, "") for c in stu_cols]
            row += [r["report_id"], r["status"]]
            if r["status"] == "graded" and r.get("grading"):
                subtotals = r["grading"].get("subtotals", {})
                row += [subtotals.get("procedural", ""), subtotals.get("semantic", ""),
                        subtotals.get("total", "")]
                scores = r["grading"].get("scores", {})
                for name, weight in dims:
                    entry = scores.get(name) or {}
                    level = entry.get("level")
                    if level in LEVEL_POINTS:
                        row.append(f"{level} {round(LEVEL_POINTS[level] * weight / 100.0, 1):g}/{weight:g}")
                    elif entry.get("score") is not None:
                        score = float(entry["score"])
                        row.append(f"{score:g} {round(score * weight / 100.0, 1):g}/{weight:g}")
                    else:
                        row.append("")
            else:
                row += ["", "", "", "N/A"] + [""] * len(dim_cols)
            ws.append(row)
        if suspicious:
            sheet = wb.create_sheet("可疑人员")
            sheet.append(stu_cols + ["报告ID", "可疑原因"])
            for s in suspicious:
                sheet.append([s["student"].get(c, "") for c in stu_cols] + [s["report_id"], s["reason"]])
        if missing:
            sheet = wb.create_sheet("缺交名单")
            sheet.append(stu_cols)
            for row_ in missing:
                sheet.append([str(row_.get(c, "") or "") for c in stu_cols])
        wb.save(path)
        return path
