"""把内容填进华东理工的空白报告模板，生成 10 份 docx、花名册与教师评分参考。

用法：.venv/Scripts/python.exe demo/build.py
"""

import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from content import EXPERIMENT, REPORTS

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "template" / "乙苯负压脱氢制苯乙烯全流程实验_空白报告.docx"

KLASS = "化工2301"
CALIBRATION = ["陈思远", "林晓桐", "王梓涵", "赵一鸣", "周雨欣"]

ROSTER = [
    ("化工2301", "化学工程学院", "2023010101", "陈思远"),
    ("化工2301", "化学工程学院", "2023010102", "林晓桐"),
    ("化工2301", "化学工程学院", "2023010103", "王梓涵"),
    ("化工2301", "化学工程学院", "2023010104", "赵一鸣"),
    ("化工2301", "化学工程学院", "2023010105", "周雨欣"),
    ("化工2301", "化学工程学院", "2023010106", "孙浩然"),
    ("化工2301", "化学工程学院", "2023010107", "吴嘉怡"),
    ("化工2301", "化学工程学院", "2023010108", "郑文博"),
    ("化工2301", "化学工程学院", "2023010109", "何佳宁"),
    ("化工2301", "化学工程学院", "2023010110", "马俊杰"),
]

TEACHER = {
    "陈思远": ["A", "A", "A", "A", "A"],
    "林晓桐": ["B", "B", "B", "C", "B"],
    "王梓涵": ["C", "C", "C", "C", "B"],
    "赵一鸣": ["D", "D", "D", "D", "C"],
    "周雨欣": ["B", "B", "B", "B", "B"],
}

SECTIONS = ["原始记录", "实验现象", "收率计算", "拓展讨论", "格式规范"]


def style(run):
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def fill(paragraph, text):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    style(paragraph.add_run(text))
    paragraph.paragraph_format.first_line_indent = None


def insert_after(anchor, text):
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    style(paragraph.add_run(text))
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = 1.5
    return paragraph


def yield_data(report):
    feed, top = report["feed"], report["top"]
    m_eb = feed * 0.867
    n_eb = m_eb / 106.17
    m_theory = n_eb * 104.15
    m_actual = top * 0.906 * 0.990
    return {"feed": feed, "top": top, "m_eb": m_eb, "n_eb": n_eb,
            "m_theory": m_theory, "m_actual": m_actual, "y": m_actual / m_theory * 100}


def yield_lines(report):
    d = yield_data(report)
    mode = report["yield_mode"]
    if mode == "volume":
        return ["以 1 h 为基准，塔顶馏出量与进料量之比即为收率：",
                f"收率 = {d['top']} / {d['feed']} × 100% = {d['top'] / d['feed'] * 100:.1f}%。"]
    head = ("以 1 h 为基准。进料为乙苯，密度取 0.867 g/mL，摩尔质量 106.17 g/mol；"
            "塔顶馏出液中苯乙烯质量分数按 99.0% 计，密度取 0.906 g/mL，摩尔质量 104.15 g/mol。")
    step1 = f"m(乙苯) = {d['feed']} mL × 0.867 g/mL = {d['m_eb']:.1f} g，n(乙苯) = {d['m_eb']:.1f} / 106.17 = {d['n_eb']:.3f} mol。"
    step2 = f"理论苯乙烯质量 = {d['n_eb']:.3f} × 104.15 = {d['m_theory']:.1f} g。"
    step3 = f"实际苯乙烯质量 = {d['top']} × 0.906 × 0.990 = {d['m_actual']:.1f} g。"
    if mode == "no_stoich":
        return [step1, step3,
                f"收率 = {d['m_actual']:.1f} / {d['m_eb']:.1f} × 100% = {d['m_actual'] / d['m_eb'] * 100:.1f}%。"]
    result = f"苯乙烯收率 = {d['m_actual']:.1f} / {d['m_theory']:.1f} × 100% = {d['y']:.1f}%。"
    if report["level"] == "C":
        return [step1, step2, step3, result]
    return [head, step1, step2, step3, result]


def build_report(report, target):
    doc = Document(TEMPLATE)
    paragraphs = doc.paragraphs

    fill(paragraphs[1], f"班　级：{KLASS}　　姓　名：{report['name']}　　学　号：{report['sid']}")
    fill(paragraphs[2], f"周　次：{report['week']}　　成　绩：＿＿＿＿＿＿")

    reaction = doc.tables[0].rows[2].cells
    for cell, value in zip(reaction, report["r101"] + report["r102"]
                           + [report["e103"], report["e104"], report["pi"]]):
        cell.text = ""
        style(cell.paragraphs[0].add_run(value))

    refine = doc.tables[1].rows[1].cells
    values = [str(report["feed"]), str(report["top"]), str(report["reflux"]),
              str(report["bottoms"]), str(report["ttop"]), str(report["tbottom"]),
              str(report["conc"]), str(report["p"])]
    for cell, value in zip(refine, values):
        cell.text = ""
        style(cell.paragraphs[0].add_run(value))

    for slot, text in zip(paragraphs[9:13], report["phenomena"]):
        fill(slot, text)
        slot.paragraph_format.first_line_indent = Pt(21)

    for slot, text in zip(paragraphs[22:27], yield_lines(report)):
        fill(slot, text)

    spare = list(paragraphs[12:19]) + list(paragraphs[34:43])
    insert_after(paragraphs[28], report["q1"])
    insert_after(paragraphs[29], report["q2"])
    insert_after(paragraphs[30], report["q3"])
    for paragraph in spare:
        paragraph._element.getparent().remove(paragraph._element)

    doc.core_properties.author = report["name"]
    doc.core_properties.title = EXPERIMENT
    doc.save(target)


def build_roster():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "花名册"
    sheet.append(["班级", "学院", "学号", "姓名"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in ROSTER:
        sheet.append(list(row))
    for column, width in zip("ABCD", (12, 18, 14, 10)):
        sheet.column_dimensions[column].width = width
    path = ROOT / "roster" / "花名册.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def build_teacher_csv():
    path = ROOT / "calibration" / "教师评分参考.csv"
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["文件名", "姓名", "学号"] + SECTIONS + ["总体"])
        for report in REPORTS:
            if report["name"] not in CALIBRATION:
                continue
            levels = TEACHER[report["name"]]
            writer.writerow([f"{report['name']}-{report['sid']}.docx", report["name"],
                             report["sid"]] + levels + [levels[0]])
    return path


def main():
    for folder in ("reports", "calibration"):
        (ROOT / folder).mkdir(exist_ok=True)
    for report in REPORTS:
        folder = "calibration" if report["name"] in CALIBRATION else "reports"
        target = ROOT / folder / f"{report['name']}-{report['sid']}.docx"
        build_report(report, target)
        print(f"{folder}/{target.name}  收率 {yield_data(report)['y']:.1f}%")
    print(build_roster())
    print(build_teacher_csv())


if __name__ == "__main__":
    main()
