"""把 cn_lab 的 20 份 md 报告转换为 docx（真实表格/标题/列表），供评分链路走 docx 提取。

用法：python tests/integration/convert_cn_lab_docx.py
md 原件移入 reports_md/ 留档，docx 输出到 reports/。
"""

import pathlib
import re
import sys

from docx import Document
from docx.shared import Pt

ROOT = pathlib.Path("tests/fixtures/bench/cn_lab")
MD_DIR = ROOT / "reports_md"
OUT_DIR = ROOT / "reports"


def convert(md_path: pathlib.Path, docx_path: pathlib.Path):
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in s.strip("|").split("|")]
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Table Grid"
            for c, text in enumerate(header):
                table.rows[0].cells[c].text = text.replace("\\(", "").replace("\\)", "")
            for r, row in enumerate(rows, start=1):
                for c, text in enumerate(row[: len(header)]):
                    table.rows[r].cells[c].text = text
            i = j
            continue
        fence = re.match(r"^(```|~~~)", s)
        if fence:
            code = []
            j = i + 1
            while j < len(lines) and not re.match(r"^(```|~~~)", lines[j].strip()):
                code.append(lines[j])
                j += 1
            p = doc.add_paragraph("\n".join(code))
            p.runs[0].font.name = "Courier New"
            i = j + 1
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", s)
        if h:
            doc.add_heading(re.sub(r"[\\$()]", "", h.group(2)).strip() or "实验报告", level=min(len(h.group(1)), 4))
            i += 1
            continue
        ul = re.match(r"^[-*+]\s+(.*)$", s)
        if ul:
            doc.add_paragraph(re.sub(r"[`*]", "", ul.group(1)), style="List Bullet")
            i += 1
            continue
        ol = re.match(r"^\d+[.、]\s+(.*)$", s)
        if ol:
            doc.add_paragraph(re.sub(r"[`*]", "", ol.group(1)), style="List Number")
            i += 1
            continue
        clean = re.sub(r"[`*]", "", s)
        if re.match(r"^【.+】", clean):
            p = doc.add_paragraph()
            run = p.add_run(clean)
            run.bold = True
        else:
            doc.add_paragraph(clean)
        i += 1
    doc.save(str(docx_path))


def main():
    MD_DIR.mkdir(exist_ok=True)
    mds = sorted(OUT_DIR.glob("*.md"))
    print("待转换:", len(mds), "份")
    for md in mds:
        target = OUT_DIR / (md.stem + ".docx")
        convert(md, target)
        md.rename(MD_DIR / md.name)
        print(" ", target.name, target.stat().st_size, "bytes")
    print("完成")


if __name__ == "__main__":
    sys.exit(main())
