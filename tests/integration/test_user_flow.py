"""端到端点击流程：上传细则 → 生成 → 保存 → 上传报告 → 提取 → 选卷 → 评分 → 终审 → 导出 → 插件。

用假模型替身跑真实路由，模拟用户在页面上的实际操作顺序。
"""
import io
import json
import shutil
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

import experiment_report_grader.factory as factory
import experiment_report_grader.web.state as state_mod
from fake_llm import FakeLLM

LLM = FakeLLM()
factory._make = lambda config: LLM

from experiment_report_grader.app import app  # noqa: E402  需在替换工厂后导入
from experiment_report_grader.web.state import state  # noqa: E402

REPORT_TEXT = """# 杨氏模量测量实验报告

## 实验数据
用逐差法处理 n0–n7 共 8 组读数，Δn_i = n_{i+4} - n_i。

## 误差分析
A 类不确定度按贝塞尔公式计算，B 类按仪器允差估算，合成得到总不确定度。
"""


def _docx(text):
    from docx import Document
    doc = Document()
    for block in text.split("\n\n"):
        doc.add_paragraph(block.replace("\n", " "))
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _wait(predicate, timeout=60):
    deadline = time.time() + timeout
    while time.time() < deadline:
        if predicate():
            return True
        time.sleep(0.05)
    raise AssertionError("等待超时")


def test_full_click_flow():
    import os
    from experiment_report_grader.mcp import store as mcp_store
    tmp = tempfile.mkdtemp(prefix="erg_flow_")
    real_output, real_sessions, real_mcp = state_mod.OUTPUT_FOLDER, state_mod.SESSIONS_DIR, mcp_store.CONFIG_PATH
    try:
        mcp_store.CONFIG_PATH = os.path.join(tmp, "mcp.yaml")
        state_mod.OUTPUT_FOLDER = tmp
        state_mod.SESSIONS_DIR = str(Path(tmp) / "sessions")
        Path(state_mod.SESSIONS_DIR).mkdir(parents=True)
        state_mod.hydrate_session(state_mod.new_session("测试课程")["id"])
        app.config["TESTING"] = True
        client = app.test_client()

        # 空课程：还没有细则，阶段应停在初始
        boot = client.get("/api/bootstrap").get_json()
        assert boot["rubric"]["count"] == 0 and boot["stage"]["key"] == "initial"

        # 1. 上传并生成细则
        resp = client.post("/upload_criteria", data={
            "criteria_file": (io.BytesIO("评分标准：误差分析 50 分，数据处理 50 分".encode()), "criteria.txt")},
            content_type="multipart/form-data")
        assert resp.status_code == 200 and resp.get_json()["filename"] == "criteria.txt"
        assert client.post("/process_criteria").status_code == 302
        _wait(lambda: (state.get("task") or {}).get("status") in ("done", "error"))
        assert state["task"]["status"] == "done", state["task"]
        assert len(state["rubric_json"]["criteria"]) == 2
        assert Path(state_mod.session_dir(), "rubrics", "rubric_current.json").exists()

        # 2. 细则页：保存编辑后的权重
        assert client.get("/app/rubric").status_code == 200
        payload = {"criteria": state["rubric_json"]["criteria"]}
        payload["criteria"][0]["weight"] = 80
        assert client.post("/rubric_table_save", json=payload).get_json()["ok"] is True
        weights = [c["weight"] for c in state["rubric_json"]["criteria"]]
        assert abs(sum(weights) - 100) < 0.01 and weights[0] > weights[1], weights

        # 3. 上传报告（含一份完全重复）
        files = {"report_files": [(_docx(REPORT_TEXT), "张三_20230001.docx"),
                                  (_docx(REPORT_TEXT), "张三_副本.docx"),
                                  (_docx(REPORT_TEXT + "\n\n补充：另做一组对照实验。"), "李四_20230002.docx")]}
        resp = client.post("/upload_reports", data=files, content_type="multipart/form-data")
        assert resp.get_json()["count"] == 3
        assert client.post("/process_reports").status_code == 302
        _wait(lambda: (state.get("task") or {}).get("status") in ("done", "error"))
        assert state["task"]["status"] == "done", state["task"]
        names = sorted(r["meta"]["filename"] for r in state["extracted_reports"])
        assert len(names) == 2 and "李四_20230002.docx" in names, names
        assert "自动去除 1 份重复" in state["task"]["detail"]

        # 4. 报告页与原文阅读器
        detail = client.get("/api/report/张三_20230001.docx").get_json()
        assert "误差分析" in json.dumps(detail, ensure_ascii=False)
        pages = client.get("/report_text/张三_20230001.docx").get_json()
        assert pages["pages"] and "逐差法" in "\n".join(pages["pages"][0])
        assert client.get("/report_original/张三_20230001.docx").status_code == 200

        # 5. 选卷 → 评分
        select = client.get("/api/select").get_json()
        assert select["hasRubric"] is True
        assert "张三_20230001.docx" in [r["filename"] for g in select["groups"]
                                       for r in g["reports"]]
        assert client.post("/pilot/confirm").status_code == 200
        resp = client.post("/start_grading", data={"report_ids": ["张三_20230001.docx"]})
        assert resp.status_code == 302 and "/app/results" in resp.headers["Location"]
        _wait(lambda: (state.get("grading_status") or {}).get("phase") == "done", 120)
        results = state["grader"].results
        assert [r["report_id"] for r in results] == ["张三_20230001.docx"]
        assert results[0]["status"] == "graded"
        assert results[0]["text_assignments"], "应把原文挂到维度"
        assert results[0]["grading"]["subtotals"]["total"] > 0

        # 6. 结果页：改等级 → 终审 → 导出
        summary = client.get("/api/results").get_json()
        assert summary["results"] and summary["counts"]["graded"] >= 1
        level_form = {"report_id": "张三_20230001.docx",
                      "level::误差分析": "B", "reason::误差分析": "证据不足"}
        client.post("/update_scores", data=level_form)
        score = state["grader"].results[0]["grading"]["scores"]["误差分析"]
        assert score["level"] == "B" and score["teacher_adjusted"] is True

        assert client.post("/approve_report/张三_20230001.docx").status_code == 302
        assert "张三_20230001.docx" in state["approved_ids"]
        assert "误差分析" in json.dumps(state["grader"].results[0]["grading"]["scores"], ensure_ascii=False)
        assert Path(state_mod.session_dir(), "approved.json").exists()

        excel = client.get("/download/grading/grading_summary.xlsx")
        assert excel.status_code == 200 and excel.data[:2] == b"PK"
        assert state["grader"]._dimension_collection("误差分析").count() >= 1

        # 7. 打回重评（带修改意见）
        resp = client.post("/regrade/张三_20230001.docx", data={"corrections": "请补充不确定度合成过程"})
        assert resp.status_code == 302
        _wait(lambda: (state.get("grading_status") or {}).get("phase") == "done", 120)
        assert "请补充不确定度合成过程" in "".join(LLM.calls)
        assert "张三_20230001.docx" not in state["approved_ids"]

        # 8. MCP 插件：登记 → 挂载 → 拉取参考 → 评分时注入
        add = client.post("/mcp/add", data={
            "name": "fake", "transport": "stdio", "command": sys.executable,
            "args": str(ROOT / "tests" / "fixtures" / "fake_mcp_server.py"), "note": "物理参考"})
        assert add.status_code == 302
        assert client.post("/mcp/mount", data={
            "servers": ["fake"], "resources": "ref://physics/*", "max_chars": "4000"}).status_code == 302
        client.post("/mcp/refresh")
        assert state["mcp_status"] and state["mcp_status"][0]["ok"] is True, state["mcp_status"]
        assert "不确定度" in state["mcp_refs"]
        client.post("/start_grading", data={"report_ids": ["张三_20230001.docx"]})
        _wait(lambda: (state.get("grading_status") or {}).get("phase") == "done", 120)
        assert any("A 类按贝塞尔公式" in call for call in LLM.calls), "插件参考应进入评分提示词"

        # 9. 课程切换后状态仍在
        client.post("/session/new", data={"name": "第二门课"})
        assert state["rubric_json"] is None and state["extracted_reports"] == []
        client.post("/session/switch", data={"session_id": state_mod.load_sessions()["sessions"][0]["id"]})
        assert state["rubric_json"] and len(state["extracted_reports"]) == 2
        assert state["grader"] and state["grader"].results
        assert state["mcp"]["servers"] == ["fake"]

        # 10. 花名册：上传 CSV → 缺交/可疑判定
        csv = "学号,姓名\n20230001,张三\n20230002,李四\n20230003,王五\n".encode("utf-8-sig")
        client.post("/roster/upload", data={"roster_file": (io.BytesIO(csv), "roster.csv")},
                    content_type="multipart/form-data")
        roster = client.get("/api/roster").get_json()["roster"]
        assert any(row.get("姓名") == "王五" for row in roster["rows"])
    finally:
        state_mod.OUTPUT_FOLDER, state_mod.SESSIONS_DIR = real_output, real_sessions
        mcp_store.CONFIG_PATH = real_mcp
        state_mod.hydrate_session(state_mod.load_sessions()["current"])
        shutil.rmtree(tmp, ignore_errors=True)


def test_pages_render_for_every_route():
    client = app.test_client()
    for url in ("/", "/app", "/app/workbench", "/app/rubric", "/task_status"):
        assert client.get(url).status_code == 200, url
    redirects = {
        "/classic": "/app/workbench",
        "/rubric_review": "/app/rubric",
        "/reports_list": "/app/reports",
        "/roster": "/app/roster",
        "/settings": "/app/settings",
        "/mcp": "/app/mcp",
        "/grading_select": "/app/select",
        "/grading_progress": "/app/results",
        "/grading_done": "/app/results",
        "/task": "/app/",
    }
    for url, target in redirects.items():
        resp = client.get(url)
        assert resp.status_code == 302 and target in resp.headers["Location"], url


if __name__ == "__main__":
    test_full_click_flow()
    print("test_user_flow ok")


def test_report_with_image_renders():
    """带插图的报告：提取图片 → 维度详情页能完整渲染并可直接取图。"""
    import os
    from PIL import Image
    from docx import Document
    from docx.shared import Inches

    tmp = tempfile.mkdtemp(prefix="erg_img_")
    real_output, real_sessions = state_mod.OUTPUT_FOLDER, state_mod.SESSIONS_DIR
    try:
        state_mod.OUTPUT_FOLDER = tmp
        state_mod.SESSIONS_DIR = str(Path(tmp) / "sessions")
        Path(state_mod.SESSIONS_DIR).mkdir(parents=True)
        state_mod.hydrate_session(state_mod.new_session("插图课")["id"])
        client = app.test_client()

        png = io.BytesIO()
        Image.new("RGB", (80, 60), "white").save(png, format="PNG")
        png.seek(0)
        doc = Document()
        doc.add_paragraph("## 装置图")
        doc.add_picture(png, width=Inches(1))
        doc.add_paragraph("图 1 实验装置示意图")
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        client.post("/upload_reports", data={"report_files": [(buffer, "带图报告.docx")]},
                    content_type="multipart/form-data")
        client.post("/process_reports")
        _wait(lambda: (state.get("task") or {}).get("status") in ("done", "error"))
        assert state["task"]["status"] == "done", state["task"]
        entry = state["extracted_reports"][0]
        assert entry["meta"]["image_count"] >= 1
        assert entry["images"], "补水后应还原图片列表"

        from urllib.parse import quote
        detail = client.get("/api/report/带图报告.docx").get_json()
        image = entry["images"][0]["filename"]
        assert image in json.dumps(detail, ensure_ascii=False)
        assert client.get(f"/report_image/{quote('带图报告.docx')}/{quote(image)}").status_code == 200
        assert client.get(f"/report_image/{quote('带图报告.docx')}/nope.png").status_code == 404
    finally:
        state_mod.OUTPUT_FOLDER, state_mod.SESSIONS_DIR = real_output, real_sessions
        state_mod.hydrate_session(state_mod.load_sessions()["current"])
        shutil.rmtree(tmp, ignore_errors=True)


def test_settings_protocol_switch():
    """配置页切换生成协议：只提交当前协议的字段，其他协议配置保持不变。"""
    import copy
    import os
    import yaml
    import experiment_report_grader.app as app_mod

    tmp = tempfile.mkdtemp(prefix="erg_cfg_")
    real_path, real_cfg = app_mod.CONFIG_PATH, app_mod.state["config"]
    try:
        app_mod.CONFIG_PATH = os.path.join(tmp, "config.yaml")
        app_mod.state["config"] = copy.deepcopy(real_cfg)
        app_mod.state["config"].setdefault("llm", {}).setdefault("openai", {})["model"] = "keep-me"
        client = app.test_client()

        resp = client.post("/settings", data={
            "llm_backend": "ollama",
            "ollama_base_url": "http://10.0.0.5:11434",
            "ollama_model": "qwen2.5:7b",
            "ocr_backend": "paddleocr",
            "vs_backend": "chroma",
            "vs_mode": "local",
        })
        assert resp.status_code == 200
        saved = yaml.safe_load(open(app_mod.CONFIG_PATH, encoding="utf-8"))
        assert saved["llm"]["backend"] == "ollama"
        assert saved["llm"]["ollama"] == {"base_url": "http://10.0.0.5:11434", "model": "qwen2.5:7b"}
        assert saved["llm"]["openai"]["model"] == "keep-me", "切换协议时不应清空其他协议配置"

        client.post("/settings", data={
            "llm_backend": "custom", "custom_base_url": "https://gw.example/v1",
            "custom_model": "gpt-x", "custom_api_key": "sk-test",
            "ocr_backend": "paddleocr", "vs_backend": "chroma", "vs_mode": "local",
        })
        saved = yaml.safe_load(open(app_mod.CONFIG_PATH, encoding="utf-8"))
        assert saved["llm"]["backend"] == "custom"
        assert saved["llm"]["custom"]["model"] == "gpt-x"
        assert saved["llm"]["custom"]["api_key"] == "sk-test"
        assert saved["llm"]["ollama"]["model"] == "qwen2.5:7b", "自定义协议不应覆盖 Ollama 配置"

        cfg_api = client.get("/api/settings").get_json()
        assert cfg_api["llm"]["backend"] == "custom"
    finally:
        app_mod.CONFIG_PATH = real_path
        app_mod.state["config"] = real_cfg
        shutil.rmtree(tmp, ignore_errors=True)


def test_spa_api_surface():
    """/app 返回 SPA 外壳，页面数据接口结构与前端约定一致。"""
    client = app.test_client()
    page = client.get("/app")
    assert page.status_code == 200
    assert b'id="root"' in page.data
    assert client.get("/favicon.ico").status_code == 200

    boot = client.get("/api/bootstrap").get_json()
    for key in ("session", "sessions", "stage", "rubric", "reports", "mcp"):
        assert key in boot

    rubric = client.get("/api/rubric").get_json()
    assert "criteria" in rubric and "hasMarkdown" in rubric

    results = client.get("/api/results").get_json()
    for key in ("results", "counts", "approvedIds", "downloads", "agreement"):
        assert key in results

    select = client.get("/api/select").get_json()
    assert "groups" in select and "hasRubric" in select

    roster = client.get("/api/roster").get_json()
    assert "roster" in roster and "suspicious" in roster

    settings = client.get("/api/settings").get_json()
    assert "llm" in settings and "vectorStore" in settings

    mcp = client.get("/api/mcp").get_json()
    assert "servers" in mcp and "mount" in mcp
