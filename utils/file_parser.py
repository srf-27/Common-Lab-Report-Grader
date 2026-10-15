import os
import struct
from pathlib import Path
from zipfile import ZipFile
import pymupdf


ARCHIVE_SIGNATURES = {
    b"PK\x03\x04": "zip",
    b"PK\x05\x06": "zip",
    b"Rar!\x1a\x07": "rar",
    b"7z\xbc\xaf\x27\x1c": "7z",
}


def detect_archive_type(filepath):
    with open(filepath, "rb") as f:
        header = f.read(8)
    for sig, fmt in ARCHIVE_SIGNATURES.items():
        if header.startswith(sig):
            return fmt
    return None


def extract_archive(filepath, dest_dir):
    fmt = detect_archive_type(filepath)
    if fmt is None:
        raise ValueError(f"Not a recognized archive: {filepath}")
    os.makedirs(dest_dir, exist_ok=True)
    if fmt == "zip":
        with ZipFile(filepath, "r") as zf:
            zf.extractall(dest_dir)
    elif fmt == "rar":
        import rarfile
        with rarfile.RarFile(filepath, "r") as rf:
            rf.extractall(dest_dir)
    elif fmt == "7z":
        import py7zr
        with py7zr.SevenZipFile(filepath, "r") as sz:
            sz.extractall(dest_dir)
    return dest_dir


def extract_pdf(filepath, output_dir, ocr_engine=None):
    os.makedirs(output_dir, exist_ok=True)
    os.makedirs(os.path.join(output_dir, "images"), exist_ok=True)

    doc = pymupdf.open(filepath)
    full_text = []
    image_records = []
    is_scanned = True
    page_count = len(doc)

    for page_num in range(page_count):
        page = doc[page_num]
        text = page.get_text("text")
        if text.strip():
            is_scanned = False
            full_text.append(f"<!-- Page {page_num + 1} -->\n{text}")

        img_index = 0
        for img_info in page.get_images(full=True):
            xref = img_info[0]
            base_image = doc.extract_image(xref)
            if base_image:
                img_bytes = base_image["image"]
                ext = base_image["ext"]
                context_label = _get_image_context(text, img_index)
                img_name = f"fig_{page_num+1:03d}_{img_index+1:03d}_{context_label}.{ext}"
                img_name = _sanitize_filename(img_name)
                img_path = os.path.join(output_dir, "images", img_name)
                with open(img_path, "wb") as f:
                    f.write(img_bytes)
                image_records.append({
                    "filename": os.path.basename(img_path),
                    "page": page_num + 1,
                    "context": context_label,
                })
                img_index += 1

    if is_scanned and ocr_engine:
        full_text = []
        for page_num in range(page_count):
            page = doc[page_num]
            pix = page.get_pixmap(dpi=300)
            img_bytes = pix.tobytes("png")
            ocr_result = ocr_engine.recognize_bytes(img_bytes)
            full_text.append(f"<!-- Page {page_num + 1} (OCR) -->\n{ocr_result}")
    elif is_scanned:
        full_text.append("WARNING: Scanned PDF. OCR engine not available.")

    text_path = os.path.join(output_dir, "text.md")
    with open(text_path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(full_text))

    doc.close()

    meta = {
        "filename": os.path.basename(filepath),
        "pages": page_count,
        "image_count": len(image_records),
        "is_scanned": is_scanned,
    }
    return meta, image_records


def extract_docx(filepath, output_dir):
    os.makedirs(output_dir, exist_ok=True)
    os.makedirs(os.path.join(output_dir, "images"), exist_ok=True)

    from docx import Document
    doc = Document(filepath)
    full_text = [p.text for p in doc.paragraphs]
    image_records = []

    docx_zip = ZipFile(filepath, "r")
    media_files = [n for n in docx_zip.namelist() if n.startswith("word/media/")]
    for idx, media_name in enumerate(media_files):
        ext = os.path.splitext(media_name)[1].lstrip(".")
        context_label = _find_docx_image_context(doc, media_name, idx)
        img_name = f"fig_{idx+1:03d}_{context_label}.{ext}"
        img_name = _sanitize_filename(img_name)
        img_path = os.path.join(output_dir, "images", img_name)
        with docx_zip.open(media_name) as src, open(img_path, "wb") as dst:
            dst.write(src.read())
        image_records.append({
            "filename": os.path.basename(img_path),
            "context": context_label,
        })
    docx_zip.close()

    text_path = os.path.join(output_dir, "text.md")
    with open(text_path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(full_text))

    meta = {
        "filename": os.path.basename(filepath),
        "paragraphs": len(doc.paragraphs),
        "image_count": len(image_records),
        "is_scanned": False,
    }
    return meta, image_records


def _get_image_context(page_text, img_index):
    lines = [l.strip() for l in page_text.split("\n") if l.strip()]
    heading_keywords = ["图", "Fig", "Figure", "表", "Table", "实验", "结果", "数据", "装置"]
    for line in reversed(lines):
        if any(kw in line for kw in heading_keywords):
            return line[:40]
    if lines:
        return lines[-1][:40]
    return "no_context"


def _find_docx_image_context(doc, media_name, img_index):
    para_texts = [p.text for p in doc.paragraphs if p.text.strip()]
    if para_texts:
        return para_texts[min(img_index, len(para_texts) - 1)][:40]
    return "no_context"


def _sanitize_filename(name):
    keepchars = (" ", ".", "_", "-")
    return "".join(c for c in name if c.isalnum() or c in keepchars).rstrip()


def collect_report_files(directory):
    supported = {".pdf", ".docx", ".zip", ".rar", ".7z"}
    files = []
    for root, _, filenames in os.walk(directory):
        for fn in filenames:
            if os.path.splitext(fn)[1].lower() in supported:
                files.append(os.path.join(root, fn))
    return files
